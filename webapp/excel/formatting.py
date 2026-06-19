from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator

import pandas as pd

from webapp.processing.constants import (
    BASELINE_CATEGORY_CURRENCY_COLUMNS,
    BASELINE_MONTHLY_CURRENCY_COLUMNS,
    BASELINE_OVERVIEW_CURRENCY_METRICS,
    EXCEL_AUTOFIT_COLUMNS,
    EXCEL_AUTOFIT_MAX_WIDTH,
    EXCEL_AUTOFIT_MIN_WIDTH,
    EXCEL_AUTOFIT_PADDING,
    EXCEL_CURRENCY_FORMAT,
    EXCEL_PERCENT_FORMAT,
    EXCEL_SHORT_DATE_FORMAT,
    HIDDEN_DESCRIPTION_COLUMNS,
    SUMMARY_DETAIL_CURRENCY_COLUMNS,
    SUMMARY_DETAIL_PERCENT_COLUMNS,
    SUMMARY_INSIGHTS_SHEET,
    SUMMARY_OVERVIEW_CURRENCY_COLUMNS,
    SUMMARY_OVERVIEW_PERCENT_COLUMNS,
)
from webapp.processing.parse import parse_amount, parse_transaction_dates


def _hide_excel_columns(worksheet, column_names: tuple[str, ...], header_row: int = 1) -> None:
    from openpyxl.utils import get_column_letter

    headers = [cell.value for cell in worksheet[header_row]]
    for col_name in column_names:
        if col_name not in headers:
            continue
        col_idx = headers.index(col_name) + 1
        worksheet.column_dimensions[get_column_letter(col_idx)].hidden = True

def _coerce_cell_currency_value(value: object) -> float | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return parse_amount(str(value))
    except (TypeError, ValueError):
        return None

def _apply_currency_columns(
    worksheet,
    column_names: tuple[str, ...],
    *,
    header_row: int = 1,
) -> None:
    """Format named columns as US dollar amounts."""
    from openpyxl.utils import get_column_letter

    headers = [cell.value for cell in worksheet[header_row]]
    for col_name in column_names:
        if col_name not in headers:
            continue
        col_idx = headers.index(col_name) + 1
        letter = get_column_letter(col_idx)
        for row in range(header_row + 1, worksheet.max_row + 1):
            cell = worksheet[f"{letter}{row}"]
            amount = _coerce_cell_currency_value(cell.value)
            if amount is None:
                continue
            cell.value = amount
            cell.number_format = EXCEL_CURRENCY_FORMAT

def _percent_column_index(headers: list[Any], col_name: str) -> int | None:
    """Match header text allowing minor spacing variants (e.g. 'vs Avg%')."""
    target = str(col_name).strip().lower().replace(" ", "")
    for idx, header in enumerate(headers):
        key = str(header or "").strip().lower().replace(" ", "")
        if key == target:
            return idx + 1
    return None

def _apply_percent_columns(
    worksheet,
    column_names: tuple[str, ...],
    *,
    header_row: int = 1,
    last_row: int | None = None,
) -> None:
    """
    Format percent columns for Excel display.

    DataFrame values are whole numbers (42.3 = 42.3%); converted to fractions for Excel.
    """
    from openpyxl.utils import get_column_letter

    headers = [cell.value for cell in worksheet[header_row]]
    end_row = last_row if last_row is not None else int(worksheet.max_row or header_row)

    for col_name in column_names:
        col_idx = _percent_column_index(headers, col_name)
        if col_idx is None:
            continue
        letter = get_column_letter(col_idx)
        for row in range(header_row + 1, end_row + 1):
            cell = worksheet[f"{letter}{row}"]
            if cell.value is None or cell.value == "":
                continue
            try:
                val = float(cell.value)
            except (TypeError, ValueError):
                continue
            if abs(val) > 1.0:
                val = val / 100.0
            cell.value = val
            cell.number_format = EXCEL_PERCENT_FORMAT

def _coerce_cell_date_value(value: object) -> datetime | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    parsed = parse_transaction_dates(pd.Series([value])).iloc[0]
    if pd.isna(parsed):
        return None
    return parsed.to_pydatetime()

def _apply_short_date_columns(
    worksheet,
    column_names: tuple[str, ...],
    *,
    header_row: int = 1,
) -> None:
    """Format named columns as Excel short date (m/d/yyyy)."""
    from openpyxl.utils import get_column_letter

    headers = [cell.value for cell in worksheet[header_row]]
    for col_name in column_names:
        if col_name not in headers:
            continue
        col_idx = headers.index(col_name) + 1
        letter = get_column_letter(col_idx)
        for row in range(header_row + 1, worksheet.max_row + 1):
            cell = worksheet[f"{letter}{row}"]
            dt = _coerce_cell_date_value(cell.value)
            if dt is None:
                continue
            cell.value = dt
            cell.number_format = EXCEL_SHORT_DATE_FORMAT

def _display_cell_length(value: object, *, number_format: str | None = None) -> int:
    if value is None:
        return 0
    if isinstance(value, float) and pd.isna(value):
        return 0
    if number_format and number_format.startswith("$") and isinstance(value, (int, float)):
        return len(f"${abs(float(value)):,.2f}") + (1 if float(value) < 0 else 0)
    return len(str(value))

def _autofit_worksheet_columns(
    worksheet,
    *,
    header_row: int = 1,
    min_width: float = EXCEL_AUTOFIT_MIN_WIDTH,
    max_width: float = EXCEL_AUTOFIT_MAX_WIDTH,
    padding: float = EXCEL_AUTOFIT_PADDING,
) -> None:
    """Approximate Excel 'Autofit Column Width' for visible columns."""
    from openpyxl.utils import get_column_letter

    if not worksheet.max_column:
        return
    last_row = max(int(worksheet.max_row or 0), header_row)
    for col_idx in range(1, worksheet.max_column + 1):
        letter = get_column_letter(col_idx)
        dim = worksheet.column_dimensions[letter]
        if dim.hidden:
            continue
        max_len = 0
        for row in range(header_row, last_row + 1):
            cell = worksheet.cell(row, col_idx)
            max_len = max(
                max_len,
                _display_cell_length(cell.value, number_format=cell.number_format),
            )
        dim.width = min(max(max_len + padding, min_width), max_width)

def autofit_workbook(writer: pd.ExcelWriter) -> None:
    """Autofit column widths on every sheet in the workbook."""
    if not EXCEL_AUTOFIT_COLUMNS:
        return
    for worksheet in writer.sheets.values():
        _autofit_worksheet_columns(worksheet)

@contextmanager
def open_excel_workbook(path: Path | str) -> Iterator[pd.ExcelWriter]:
    """
    Context manager for Excel output. All sheets get column autofit on save (default).
    """
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        yield writer
        autofit_workbook(writer)

def _enable_worksheet_autofilter(worksheet) -> None:
    """Turn on Excel column filters on row 1 (Filter button in header)."""
    from openpyxl.utils import get_column_letter

    max_row = int(worksheet.max_row or 0)
    max_col = int(worksheet.max_column or 0)
    if max_row < 1 or max_col < 1:
        return
    worksheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

def format_transaction_worksheet(
    worksheet,
    *,
    hide_descriptions: bool = False,
    short_date_columns: tuple[str, ...] = (),
    enable_autofilter: bool = False,
) -> None:
    if hide_descriptions:
        _hide_excel_columns(worksheet, HIDDEN_DESCRIPTION_COLUMNS)
    if short_date_columns:
        _apply_short_date_columns(worksheet, short_date_columns)
    _apply_currency_columns(worksheet, ("Amount",))
    worksheet.freeze_panes = "A2"
    if enable_autofilter:
        _enable_worksheet_autofilter(worksheet)

def format_baseline_worksheet(
    worksheet,
    *,
    overview_rows: int,
    categories_header_row: int,
    monthly_header_row: int | None = None,
) -> None:
    """Apply currency formats to the three Baseline tables."""
    headers = [cell.value for cell in worksheet[1]]
    if "Metric" in headers and "Value" in headers:
        metric_col = headers.index("Metric") + 1
        value_col = headers.index("Value") + 1
        for row in range(2, overview_rows + 2):
            metric = worksheet.cell(row, metric_col).value
            if metric not in BASELINE_OVERVIEW_CURRENCY_METRICS:
                continue
            cell = worksheet.cell(row, value_col)
            amount = _coerce_cell_currency_value(cell.value)
            if amount is not None:
                cell.value = amount
                cell.number_format = EXCEL_CURRENCY_FORMAT

    _apply_currency_columns(
        worksheet, BASELINE_CATEGORY_CURRENCY_COLUMNS, header_row=categories_header_row
    )
    if monthly_header_row is not None:
        _apply_currency_columns(
            worksheet, BASELINE_MONTHLY_CURRENCY_COLUMNS, header_row=monthly_header_row
        )

def format_spending_insights_worksheet(
    worksheet,
    *,
    overview_rows: int,
    overview_header_row: int,
    detail_header_rows: list[int],
    month_anchors: dict[str, int],
) -> None:
    """Currency formats and 'View categories' hyperlinks on the Summary insights sheet."""
    from openpyxl.styles import Font

    overview_last_row = overview_header_row + overview_rows
    _apply_currency_columns(
        worksheet,
        SUMMARY_OVERVIEW_CURRENCY_COLUMNS,
        header_row=overview_header_row,
    )
    _apply_percent_columns(
        worksheet,
        SUMMARY_OVERVIEW_PERCENT_COLUMNS,
        header_row=overview_header_row,
        last_row=overview_last_row,
    )

    sorted_detail_headers = sorted(detail_header_rows)
    for i, header_row in enumerate(sorted_detail_headers):
        if i + 1 < len(sorted_detail_headers):
            block_end = sorted_detail_headers[i + 1] - 2
        else:
            block_end = int(worksheet.max_row or header_row)
        _apply_currency_columns(
            worksheet,
            SUMMARY_DETAIL_CURRENCY_COLUMNS,
            header_row=header_row,
        )
        _apply_percent_columns(
            worksheet,
            SUMMARY_DETAIL_PERCENT_COLUMNS,
            header_row=header_row,
            last_row=block_end,
        )

    headers = [cell.value for cell in worksheet[overview_header_row]]
    if "View Categories" not in headers:
        return
    link_col = headers.index("View Categories") + 1
    month_col = headers.index("Budget Month") + 1 if "Budget Month" in headers else None
    link_font = Font(color="0563C1", underline="single")

    for row in range(overview_header_row + 1, overview_header_row + overview_rows + 1):
        month_val = (
            worksheet.cell(row, month_col).value if month_col is not None else None
        )
        month_key = str(month_val or "").strip()
        anchor = month_anchors.get(month_key)
        if not anchor:
            continue
        cell = worksheet.cell(row, link_col)
        cell.value = "View categories"
        cell.hyperlink = f"#'{SUMMARY_INSIGHTS_SHEET}'!A{anchor}"
        cell.font = link_font
