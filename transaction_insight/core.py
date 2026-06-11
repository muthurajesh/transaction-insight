#!/usr/bin/env python3
"""
Process raw bank transaction CSV exports into a structured Excel workbook
with AI-assisted categorization and fixed/variable cost classification.
"""

from __future__ import annotations

import argparse
import calendar
import hashlib
import json
import os
import re
import sys
import time
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterator

import pandas as pd
from dotenv import load_dotenv
from openai import OpenAI

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
CONFIG_DIR = PROJECT_ROOT / "config"
INPUT_DIR = PROJECT_ROOT / "input"
OUTPUT_DIR = PROJECT_ROOT / "output"

load_dotenv(CONFIG_DIR / ".env")

BATCH_SIZE = int(os.getenv("BATCH_SIZE", "20"))
LOCAL_BATCH_SIZE = int(os.getenv("LOCAL_BATCH_SIZE", "40"))
DESCRIPTION_BATCH_SIZE = int(os.getenv("DESCRIPTION_BATCH_SIZE", "0"))
CLASSIFICATION_BATCH_SIZE = int(os.getenv("CLASSIFICATION_BATCH_SIZE", "0"))
LM_STUDIO_BASE_URL = os.getenv("LM_STUDIO_BASE_URL", "http://192.168.0.7:1234/v1")
LM_STUDIO_MODEL = os.getenv("LM_STUDIO_MODEL", "qwen2.5-coder-32b-instruct")
OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434/v1")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:14b")
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
REQUEST_TIMEOUT = float(os.getenv("REQUEST_TIMEOUT", "600"))
PAYROLL_SPILLOVER_DAYS = int(os.getenv("PAYROLL_SPILLOVER_DAYS", "7"))
LOOKUP_FILENAME = os.getenv("LOOKUP_FILE", "transaction-lookups.xlsx")
HISTORY_FILENAME = os.getenv("HISTORY_FILE", "transaction-history.xlsx")
BASELINE_VARIABLE_BUFFER_PCT = float(os.getenv("BASELINE_VARIABLE_BUFFER_PCT", "10"))
BASELINE_FIXED_BUFFER_PCT = float(os.getenv("BASELINE_FIXED_BUFFER_PCT", "5"))
BASELINE_RECOMMENDED_MONTHS = int(os.getenv("BASELINE_RECOMMENDED_MONTHS", "3"))
CADENCE_SPIKE_RATIO = float(os.getenv("CADENCE_SPIKE_RATIO", "2.0"))
CADENCE_MIN_HISTORY_MONTHS = int(os.getenv("CADENCE_MIN_HISTORY_MONTHS", "3"))
CADENCE_YEARLY_GAP_MIN = int(os.getenv("CADENCE_YEARLY_GAP_MIN", "10"))
CADENCE_YEARLY_GAP_MAX = int(os.getenv("CADENCE_YEARLY_GAP_MAX", "14"))
LOCAL_PROVIDERS = frozenset({"lmstudio", "ollama"})

DROP_OUTPUT_COLUMNS = ("Status", "Split Type", "Currency", "Memo")
HIDDEN_DESCRIPTION_COLUMNS = ("Original Description", "User Description", "Simple Description")
TRANSACTION_SHEETS = ("Raw Data", "Income", "Expenses", "Adjustments")
DETAIL_TRANSACTION_SHEETS = ("Income", "Expenses", "Adjustments")

EXCEL_CURRENCY_FORMAT = "$#,##0.00"
EXCEL_PERCENT_FORMAT = "0.0%"
EXCEL_SHORT_DATE_FORMAT = "m/d/yyyy"
EXCEL_AUTOFIT_COLUMNS = True  # applied to every sheet on all workbooks this script writes
EXCEL_AUTOFIT_MIN_WIDTH = 8
EXCEL_AUTOFIT_MAX_WIDTH = 55
EXCEL_AUTOFIT_PADDING = 1.5

SUMMARY_INSIGHTS_SHEET = "Summary"
SUMMARY_OVERVIEW_CURRENCY_COLUMNS = (
    "Gross Income",
    "Total Expenses",
    "Avg Monthly Expenses",
    "vs Avg",
)
SUMMARY_OVERVIEW_PERCENT_COLUMNS = ("vs Avg %",)
SUMMARY_DETAIL_CURRENCY_COLUMNS = (
    "Month Spend",
    "Avg Monthly Spend",
    "vs Avg",
)
SUMMARY_DETAIL_PERCENT_COLUMNS = ("vs Avg %", "% of Month")

BASELINE_CATEGORY_CURRENCY_COLUMNS = (
    "Total Spend",
    "Typical Monthly Spend (Core)",
    "Avg Monthly Spend",
    "Min Month Spend",
    "Max Month Spend",
    "Suggested Monthly Budget",
)

BASELINE_MONTHLY_CURRENCY_COLUMNS = (
    "Total Spend",
    "Core Monthly Spend",
    "vs Monthly Avg",
)

BASELINE_OVERVIEW_CURRENCY_METRICS = frozenset(
    {
        "Overall avg monthly spend",
        "Overall typical core monthly spend",
        "Sum of suggested category budgets",
    }
)

LOOKUP_SHEETS = (
    "Categories",
    "CategoryRules",
    "Types",
    "BusinessCategoryRules",
    "DescriptionLookup",
    "MerchantCategories",
)

MERCHANT_CATEGORIES_SHEET = "MerchantCategories"
MERCHANT_CATEGORY_COLUMNS = (
    "Merchant Key",
    "AI Category",
    "AI Sub-Category",
    "Budget Tier",
    "Type",
    "Flow Type",
    "Classification",
    "Transaction Count",
    "Notes",
)

EXPENSE_CADENCE_RULES_SHEET = "ExpenseCadenceRules"
EXPENSE_CADENCE_RULE_COLUMNS = (
    "Generated Description",
    "Cadence",
    "In Monthly Run-Rate?",
    "Notes",
)
CADENCE_MONTHLY = "Monthly"
CADENCE_YEARLY = "Yearly"
CADENCE_ONETIME = "One-time"
CADENCE_UNPLANNED = "Unplanned"
CADENCE_UNKNOWN = "Unknown"
CADENCE_VALID = frozenset(
    {CADENCE_MONTHLY, CADENCE_YEARLY, CADENCE_ONETIME, CADENCE_UNPLANNED, CADENCE_UNKNOWN}
)
CADENCE_SOURCE_LOOKUP = "Lookup"
CADENCE_SOURCE_DETECTED = "Detected"
CADENCE_SOURCE_DEFAULT = "Default"
CADENCE_DEFAULT_RUNRATE: dict[str, str] = {
    CADENCE_MONTHLY: "Y",
    CADENCE_YEARLY: "N",
    CADENCE_ONETIME: "N",
    CADENCE_UNPLANNED: "N",
    CADENCE_UNKNOWN: "Y",
}
CADENCE_REVIEW_SHEET = "Cadence Review"

CUSTOM_RULES_SHEET = "CustomRules"
CUSTOM_RULES_COLUMNS = ("Rule", "Status", "Compiled Rule", "Last Error", "Updated At")
CUSTOM_RULE_STATUS_PENDING = "Pending"
CUSTOM_RULE_STATUS_ACTIVE = "Active"
CUSTOM_RULE_STATUS_ERROR = "Error"
CUSTOM_RULE_STATUS_DISABLED = "Disabled"

CUSTOM_RULE_FIELD_MAP = {
    "category": "Category",
    "classification": "Classification",
    "ai_category": "AI Category",
    "ai_sub_category": "AI Sub-Category",
    "type": "Type",
    "sub_type": "Sub-Type",
    "budget_tier": "Budget Tier",
}

CUSTOM_RULE_COMPILER_PROMPT = """You translate plain-English transaction rules into JSON for an automated processor.

Return ONLY valid JSON: {"rule": { ... }}

Supported rule_type values:

1. "assign" — set fields when match criteria all apply (AND). All text matching is case-insensitive.
   Use *text* for contains, prefix* for starts-with, *suffix for ends-with; plain text is exact match.
   Optional amount: compare absolute dollar value (9.99 matches -9.99 and 9.99).
   Use match.description when the user says "description" (searches Generated, Original, Simple, and User descriptions).
   Use match.generated_description only for the merchant/payee label (Generated Description).
   For OR conditions, use a JSON array of patterns, e.g. ["*amazon web services*","*aws*"].
   {"rule_type":"assign","match":{"description":["*amazon web services*","*aws*"]},"set":{"ai_category":"Business Expenses","ai_sub_category":"Cloud Computing/Hosting","type":"Variable","classification":"Business"}}
   {"rule_type":"assign","match":{"generated_description":"*check*","amount":"60"},"set":{"ai_category":"Education","ai_sub_category":"Music Lessons"}}
   {"rule_type":"assign","match":{"generated_description":"Apple","amount":"9.99"},"set":{"category":"Business Expenses","classification":"Business","ai_category":"Business","ai_sub_category":"Laptop Warranty"}}

2. "monthly_split_max" — rows with the same Generated Description in the same calendar/budget month:
   the row with the largest absolute Amount gets when_max; every other row in that month gets when_other.
   Use for "multiple entries per month, highest is X, others are Y".
   {"rule_type":"monthly_split_max","match":{"generated_description":"Ahs Ahs.Com"},"group_by":"Budget Month","min_rows_per_group":2,"when_max":{"category":"Insurance","ai_category":"Utilities","ai_sub_category":"Appliance Insurance"},"when_other":{"category":"Business Expenses","classification":"Business","ai_category":"Rental","ai_sub_category":"Appliance Insurance"}}

Allowed field keys in set / when_max / when_other: category, classification, ai_category, ai_sub_category, type, sub_type, budget_tier.
Use exact Generated Description spelling from the user's rule when possible."""

DESCRIPTION_LOOKUP_COLUMNS = (
    "Source Key",
    "User Description",
    "Simple Description",
    "Original Description",
    "Generated Description",
    "Source",
    "Model",
    "Updated At",
)

HISTORY_SHEETS = ("Income", "Expenses", "Adjustments", "Summary", "Baseline")

DESCRIPTION_PROMPT = """You are a personal finance assistant. For each transaction, read the bank's
Original Description, optional User Description, and Simple Description. Produce one short
Generated Description: a clear merchant or payee label (3–8 words) suitable for budgeting reports.
Prefer Simple Description or User Description when they name the merchant; otherwise distill
Original Description (ignore card numbers, DES:/ID:/INDN:/CO ID: boilerplate).
Return ONLY valid JSON: {"results": [{"index": <int>, "generated_description": "<string>"}]}."""

BUSINESS_RULE_PROMPT = """You are a bookkeeper. Each transaction is marked Business (not Personal).
Suggest how future similar charges should be categorized for a personal+business export workflow.
Return ONLY valid JSON: {"results": [{"index": <int>, "ai_category": "<string>",
"ai_sub_category": "<string>", "budget_tier": "Need|Want|Wish", "type": "Fixed|Variable",
"notes": "<short rationale>"}]}."""

CLASSIFICATION_PROMPT = """You are a personal finance analyst. Classify each transaction below.

For each transaction, return:
1. section: "Income" or "Expense"
   - Income: salary, interest, dividends, refunds/credits that reduce prior spending, legitimate incoming money
   - Expense: money spent or outflows (including transfers out, bill payments, purchases)
   - Credit card "payment received" / "online payment thank you" on a credit card account is NOT income — treat as Expense (internal transfer)
   - Zelle/Venmo received from individuals: use context — small peer payments may be Expense-related reimbursements; payroll-like amounts are Income
2. category: clear top-level category (e.g. Housing, Utilities, Groceries, Dining, Transportation, Healthcare, Insurance, Entertainment, Income, Transfers, Savings, Subscriptions, Pets, Shopping, Personal Care, Education, Charitable, Fees, Other)
3. sub_category: short bill/spend type (e.g. "Towing", "Electric bill", "Fast food", "Fuel") — NOT the merchant or store name (that is already in description)
4. type: "Fixed" or "Variable"
   - Fixed: recurring obligations you expect each month even if the amount varies slightly (mortgage, rent, HOA, utilities, insurance premiums, phone/internet, subscriptions, gym membership, loan payments, childcare, minimum debt payments)
   - Variable: discretionary or fluctuating spending (groceries, restaurants, fuel, shopping, entertainment outings, gifts, one-off purchases)
   - Nuance: electricity/gas/water bills are Fixed (required monthly). Groceries and dining are Variable.
5. sub_type: always return empty string "" (reserved for future use)

Return ONLY valid JSON: an array of objects with keys: index, section, category, sub_category, type, sub_type.
The index must match the transaction index provided."""


def json_for_prompt(obj: Any) -> str:
    """Compact JSON for LLM prompts (smaller/faster than pretty-printed)."""
    return json.dumps(obj, separators=(",", ":"))


def resolve_batch_sizes(
    provider: str,
    *,
    batch_size_arg: int | None,
) -> tuple[int, int, int]:
    """
    Return (general, description, classification) batch sizes.
    Local providers default to LOCAL_BATCH_SIZE unless --batch-size is set.
    """
    if batch_size_arg is not None:
        base = batch_size_arg
    elif provider in LOCAL_PROVIDERS:
        base = LOCAL_BATCH_SIZE
    else:
        base = BATCH_SIZE

    desc = DESCRIPTION_BATCH_SIZE if DESCRIPTION_BATCH_SIZE > 0 else base
    if CLASSIFICATION_BATCH_SIZE > 0:
        classify = CLASSIFICATION_BATCH_SIZE
    else:
        # Classification returns more tokens per row; cap slightly on local runs.
        classify = min(base, 25) if provider in LOCAL_PROVIDERS else base
    return base, desc, classify


def format_duration(seconds: float) -> str:
    """Human-readable duration for console output."""
    if seconds < 0.001:
        return "<1 ms"
    if seconds < 1:
        return f"{seconds * 1000:.0f} ms"
    if seconds < 60:
        return f"{seconds:.2f} s"
    minutes, secs = divmod(seconds, 60)
    if minutes < 60:
        return f"{int(minutes)}m {secs:.1f}s"
    hours, minutes = divmod(minutes, 60)
    return f"{int(hours)}h {int(minutes)}m {secs:.0f}s"


class PhaseTimer:
    """Collect per-phase wall times and print a summary at the end of a run."""

    def __init__(self) -> None:
        self._run_start = time.perf_counter()
        self.phases: list[tuple[str, float]] = []

    @contextmanager
    def phase(self, name: str) -> Iterator[None]:
        t0 = time.perf_counter()
        yield
        self.phases.append((name, time.perf_counter() - t0))

    @staticmethod
    @contextmanager
    def track(timer: PhaseTimer | None, name: str) -> Iterator[None]:
        """Record a named phase when *timer* is set; no-op otherwise."""
        if timer is None:
            yield
        else:
            with timer.phase(name):
                yield

    def total_seconds(self) -> float:
        return time.perf_counter() - self._run_start

    def print_summary(self) -> None:
        total = self.total_seconds()
        phase_sum = sum(d for _, d in self.phases)
        llm_phase_names = {"Descriptions (LLM)", "AI classification (LLM)"}
        print("\nTiming summary")
        print("─" * 52)
        name_width = max((len(n) for n, _ in self.phases), default=20)
        for name, elapsed in self.phases:
            if elapsed < 0.05 and name not in llm_phase_names:
                continue
            pct = (elapsed / total * 100) if total > 0 else 0
            print(
                f"  {name:<{name_width}}  {format_duration(elapsed):>10}  ({pct:4.1f}%)"
            )
        llm_total = sum(d for n, d in self.phases if n in llm_phase_names)
        if llm_total > 0 and len(llm_phase_names.intersection(n for n, _ in self.phases)) > 1:
            pct = (llm_total / total * 100) if total > 0 else 0
            print("─" * 52)
            print(
                f"  {'LLM total':<{name_width}}  {format_duration(llm_total):>10}  ({pct:4.1f}%)"
            )
        print("─" * 52)
        print(f"  {'Total':<{name_width}}  {format_duration(total):>10}  (100.0%)")
        other = total - phase_sum
        if other > 0.05:
            print(
                f"  (Phases account for {format_duration(phase_sum)}; "
                f"setup/print overhead {format_duration(other)})"
            )


def parse_amount(value: str | float) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = str(value).strip().replace(",", "").replace("$", "")
    return float(cleaned) if cleaned else 0.0


def parse_transaction_dates(date_series: pd.Series) -> pd.Series:
    """
    Parse bank export dates into datetime64.

    Supports common US export formats (M/D/Y and M/D/YY), ISO dates, and values
    already stored as datetimes (e.g. from Excel history sheets).
    """
    if date_series.empty:
        return pd.Series(dtype="datetime64[ns]")

    if pd.api.types.is_datetime64_any_dtype(date_series):
        return pd.to_datetime(date_series, errors="coerce")

    raw = date_series.fillna("").astype(str).str.strip()
    empty_mask = raw.isin(("", "nan", "NaT", "None", "nat"))
    result = pd.Series(pd.NaT, index=date_series.index, dtype="datetime64[ns]")

    if not empty_mask.all():
        values = raw[~empty_mask]
        # pandas 2+: mixed M/D/Y, M/D/YY, ISO, etc. in one column
        try:
            mixed = pd.to_datetime(values, format="mixed", errors="coerce")
            result.loc[mixed.index] = mixed
        except (ValueError, TypeError):
            pass

        still_missing = result.isna() & ~empty_mask
        explicit_formats = (
            "%m/%d/%y",
            "%m/%d/%Y",
            "%m-%d-%y",
            "%m-%d-%Y",
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%Y-%m-%d %H:%M:%S",
            "%m/%d/%Y %H:%M:%S",
            "%m/%d/%y %H:%M:%S",
            "%d-%b-%Y",
            "%d-%b-%y",
            "%b %d, %Y",
            "%b %d, %y",
        )
        for fmt in explicit_formats:
            if not still_missing.any():
                break
            attempt = pd.to_datetime(
                raw.loc[still_missing], format=fmt, errors="coerce"
            )
            result.loc[still_missing] = attempt
            still_missing = result.isna() & ~empty_mask

        if still_missing.any():
            inferred = pd.to_datetime(
                raw.loc[still_missing], errors="coerce", dayfirst=False
            )
            result.loc[still_missing] = inferred

    return result


def drop_rows_with_invalid_dates(
    df: pd.DataFrame,
    *,
    date_column: str = "Date",
    out_column: str = "Transaction Date",
) -> tuple[pd.DataFrame, int]:
    """Parse transaction dates; drop rows that cannot be parsed. Returns (df, dropped_count)."""
    if date_column not in df.columns:
        raise ValueError(f"Missing column: {date_column}")

    n_before = len(df)
    out = df.copy()
    out[out_column] = parse_transaction_dates(out[date_column])
    bad_mask = out[out_column].isna()
    dropped = int(bad_mask.sum())
    if dropped:
        examples = out.loc[bad_mask, date_column].astype(str).head(5).tolist()
        print(
            f"  Warning: dropped {dropped} of {n_before} row(s) with unparseable "
            f"'{date_column}'",
            flush=True,
        )
        if examples:
            print(f"  Unparseable examples: {examples}", flush=True)
    out = out.loc[~bad_mask].copy()
    return out, dropped


def merchant_key(row: pd.Series) -> str:
    """Best-effort merchant identifier for lookup matching."""
    explicit = str(row.get("Merchant Key", "") or "").strip()
    if explicit:
        return explicit[:120]
    gd = str(row.get("Generated Description", "") or "").strip()
    if gd:
        return " ".join(gd.split())[:120]
    sd = str(row.get("Simple Description", "") or "").strip()
    od = str(row.get("Original Description", "") or "").strip()
    key = sd if sd else od
    return " ".join(key.split())[:120]


def ensure_merchant_key_column(df: pd.DataFrame) -> pd.DataFrame:
    """Populate Merchant Key from generated description when missing."""
    if "Merchant Key" not in df.columns:
        df = df.copy()
        df["Merchant Key"] = df.apply(merchant_key, axis=1)
        return df
    missing = df["Merchant Key"].fillna("").astype(str).str.strip() == ""
    if missing.any():
        df = df.copy()
        df.loc[missing, "Merchant Key"] = df.loc[missing].apply(merchant_key, axis=1)
    return df


def _row_merchant_key(row: pd.Series) -> str:
    return str(row.get("Merchant Key", "") or merchant_key(row)).strip()


def _is_semantic_sub_category(row: pd.Series) -> bool:
    sub = str(row.get("AI Sub-Category", "") or "").strip()
    if not sub:
        return False
    mk = _row_merchant_key(row).lower()
    return sub.lower() != mk


def _merchant_category_is_semantic(row: pd.Series) -> bool:
    mk = str(row.get("Merchant Key", "") or "").strip().lower()
    sub = str(row.get("AI Sub-Category", "") or "").strip().lower()
    return bool(sub) and sub != mk


def is_business_row(row: pd.Series) -> bool:
    return str(row.get("Classification", "") or "").strip().lower() == "business"


def _clean_original_for_display(text: str) -> str:
    """Strip common bank noise from Original Description for heuristic labels."""
    t = " ".join(str(text or "").split())
    if not t:
        return ""
    t = re.sub(r"\bx{6,}\d+\b", "", t, flags=re.I)
    t = re.sub(r"\bDES:.*", "", t, flags=re.I)
    t = re.sub(r"\bID:.*", "", t, flags=re.I)
    t = re.sub(r"\bINDN:.*", "", t, flags=re.I)
    t = re.sub(r"\bCO ID:.*", "", t, flags=re.I)
    t = re.sub(r"\bMOBILE PURCHASE\s+\d{4}\s*", "", t, flags=re.I)
    return " ".join(t.split()).strip()[:120]


def heuristic_generated_description(row: pd.Series) -> str:
    user = str(row.get("User Description", "") or "").strip()
    simple = str(row.get("Simple Description", "") or "").strip()
    original = _clean_original_for_display(str(row.get("Original Description", "") or ""))
    if user:
        return user[:120]
    if simple:
        return simple[:120]
    return original[:120] if original else "Unknown"


def _norm_description_part(text: str) -> str:
    return " ".join(str(text or "").split()).strip().lower()


def description_source_key(row: pd.Series) -> str:
    """Stable hash key from normalized User, Simple, and cleaned Original descriptions."""
    user = _norm_description_part(row.get("User Description", ""))
    simple = _norm_description_part(row.get("Simple Description", ""))
    original = _norm_description_part(
        _clean_original_for_display(str(row.get("Original Description", "") or ""))
    )
    payload = f"u:{user}|s:{simple}|o:{original[:300]}"
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


def build_description_lookup_map(lookups: dict[str, pd.DataFrame]) -> dict[str, str]:
    """Source Key -> Generated Description from the DescriptionLookup sheet."""
    sheet = lookups.get("DescriptionLookup")
    if sheet is None or sheet.empty:
        return {}
    if "Source Key" not in sheet.columns or "Generated Description" not in sheet.columns:
        return {}
    result: dict[str, str] = {}
    for _, row in sheet.iterrows():
        key = str(row.get("Source Key", "") or "").strip()
        desc = str(row.get("Generated Description", "") or "").strip()
        if key and desc:
            result[key] = desc[:120]
    return result


def make_description_lookup_row(
    row: pd.Series,
    source_key: str,
    generated: str,
    source: str,
    model: str,
) -> dict[str, str]:
    return {
        "Source Key": source_key,
        "User Description": str(row.get("User Description", "") or "")[:200],
        "Simple Description": str(row.get("Simple Description", "") or "")[:200],
        "Original Description": str(row.get("Original Description", "") or "")[:300],
        "Generated Description": generated[:120],
        "Source": source,
        "Model": model if source == "llm" else "",
        "Updated At": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def merge_description_lookup(
    existing: pd.DataFrame | None,
    new_rows: pd.DataFrame | None,
    *,
    rebuild: bool = False,
) -> pd.DataFrame:
    """Merge description lookup rows. Default: append new keys only. Rebuild: upsert by Source Key."""
    cols = list(DESCRIPTION_LOOKUP_COLUMNS)
    if new_rows is None or new_rows.empty:
        if existing is not None and not existing.empty:
            out = existing.copy()
            for col in cols:
                if col not in out.columns:
                    out[col] = ""
            return out[cols].sort_values("Source Key").reset_index(drop=True)
        return pd.DataFrame(columns=cols)

    new_rows = new_rows.copy()
    for col in cols:
        if col not in new_rows.columns:
            new_rows[col] = ""

    if existing is None or existing.empty:
        return (
            new_rows[cols]
            .drop_duplicates(subset=["Source Key"], keep="last")
            .sort_values("Source Key")
            .reset_index(drop=True)
        )

    existing = existing.copy()
    for col in cols:
        if col not in existing.columns:
            existing[col] = ""

    if rebuild:
        merged = pd.concat([existing[cols], new_rows[cols]], ignore_index=True)
        return (
            merged.drop_duplicates(subset=["Source Key"], keep="last")
            .sort_values("Source Key")
            .reset_index(drop=True)
        )

    existing_keys = set(existing["Source Key"].fillna("").astype(str).str.strip())
    append = new_rows[
        ~new_rows["Source Key"].fillna("").astype(str).str.strip().isin(existing_keys)
    ]
    merged = pd.concat([existing[cols], append[cols]], ignore_index=True)
    return merged.sort_values("Source Key").reset_index(drop=True)


def build_description_payload(row: pd.Series, index: int) -> dict[str, Any]:
    return {
        "index": index,
        "original_description": str(row.get("Original Description", "") or "")[:300],
        "user_description": str(row.get("User Description", "") or "")[:200],
        "simple_description": str(row.get("Simple Description", "") or "")[:200],
        "category": str(row.get("Category", "") or ""),
        "amount": str(row.get("Amount", "") or ""),
    }


def generate_descriptions_batch(
    client: OpenAI,
    transactions: list[dict[str, Any]],
    model: str,
    *,
    use_json_mode: bool,
) -> list[dict[str, Any]]:
    user_content = json_for_prompt(transactions)
    messages = [
        {"role": "system", "content": DESCRIPTION_PROMPT},
        {
            "role": "user",
            "content": (
                "Generate descriptions for these transactions. Respond with JSON "
                '{"results": [...]} only.\n\n'
                f"Transactions:\n{user_content}"
            ),
        },
    ]
    kwargs: dict[str, Any] = {
        "model": model,
        "temperature": 0.1,
        "messages": messages,
    }
    if use_json_mode:
        kwargs["response_format"] = {"type": "json_object"}

    response = client.chat.completions.create(**kwargs)
    raw = response.choices[0].message.content or "{}"
    parsed = extract_json_payload(raw)
    results = parsed.get("results", parsed if isinstance(parsed, list) else [])
    if not isinstance(results, list):
        raise ValueError(f"Unexpected description response: {raw[:500]}")
    return results


def fill_generated_descriptions(
    df: pd.DataFrame,
    client: OpenAI,
    model: str,
    *,
    batch_size: int,
    use_json_mode: bool,
    description_lookup: dict[str, str] | None = None,
    rebuild_lookup: bool = False,
    on_batch_progress: Callable[[int, int, str], None] | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """
    Set Generated Description using shared lookup, User Description, then LLM.
    When rebuild_lookup is True, cache is ignored and LLM runs for all non-user rows.
    Returns (dataframe, new lookup rows to merge into transaction-lookups.xlsx).
    """
    df = df.copy()
    lookup = {} if rebuild_lookup else dict(description_lookup or {})
    new_entries: list[dict[str, str]] = []
    llm_indices: list[int] = []
    lookup_hits = 0
    user_hits = 0

    for idx, row in df.iterrows():
        key = description_source_key(row)
        cached = lookup.get(key, "")
        if cached and not rebuild_lookup:
            df.at[idx, "Generated Description"] = cached
            lookup_hits += 1
            continue

        user = str(row.get("User Description", "") or "").strip()
        if user:
            desc = user[:120]
            df.at[idx, "Generated Description"] = desc
            new_entries.append(make_description_lookup_row(row, key, desc, "user", model))
            lookup[key] = desc
            user_hits += 1
            continue

        llm_indices.append(int(idx))

    if rebuild_lookup:
        print(
            f"  Rebuilding description lookup: {user_hits} from User Description, "
            f"{len(llm_indices)} row(s) need LLM",
            flush=True,
        )
    else:
        print(
            f"  Description lookup: {lookup_hits} hit(s), {user_hits} from User Description, "
            f"{len(llm_indices)} row(s) need LLM",
            flush=True,
        )

    for idx in llm_indices:
        df.at[idx, "Generated Description"] = heuristic_generated_description(df.loc[idx])

    if llm_indices:
        key_to_indices: dict[str, list[int]] = {}
        for idx in llm_indices:
            key = description_source_key(df.loc[idx])
            key_to_indices.setdefault(key, []).append(idx)

        llm_tasks = [(key, indices[0]) for key, indices in key_to_indices.items()]
        if len(llm_tasks) < len(llm_indices):
            print(
                f"  Deduped description LLM: {len(llm_indices)} rows -> "
                f"{len(llm_tasks)} unique pattern(s)",
                flush=True,
            )

        payloads = [
            build_description_payload(df.loc[rep_idx], batch_idx)
            for batch_idx, (_, rep_idx) in enumerate(llm_tasks)
        ]
        total_payloads = len(payloads)
        for start in range(0, total_payloads, batch_size):
            batch = payloads[start : start + batch_size]
            end = start + len(batch)
            batch_msg = f"Descriptions (LLM) {start + 1}–{end} of {total_payloads}…"
            print(f"  {batch_msg}", flush=True)
            if on_batch_progress:
                on_batch_progress(
                    start,
                    total_payloads,
                    f"{batch_msg} (calling model)",
                )
            try:
                batch_results = generate_descriptions_batch(
                    client, batch, model, use_json_mode=use_json_mode
                )
            except Exception as exc:
                print(f"  Description LLM failed; keeping heuristics: {exc}", flush=True)
                break
            if on_batch_progress:
                on_batch_progress(end, total_payloads, batch_msg)
            for item in batch_results:
                batch_idx = int(item.get("index", -1))
                desc = str(item.get("generated_description", "") or "").strip()
                if batch_idx < 0 or batch_idx >= len(llm_tasks) or not desc:
                    continue
                key, rep_idx = llm_tasks[batch_idx]
                for idx in key_to_indices[key]:
                    df.at[idx, "Generated Description"] = desc[:120]
                new_entries.append(
                    make_description_lookup_row(
                        df.loc[rep_idx], key, desc, "llm", model
                    )
                )
                lookup[key] = desc[:120]

    new_lookup_df = (
        pd.DataFrame(new_entries, columns=list(DESCRIPTION_LOOKUP_COLUMNS))
        if new_entries
        else pd.DataFrame(columns=list(DESCRIPTION_LOOKUP_COLUMNS))
    )
    return df, new_lookup_df


def transaction_fingerprint(row: pd.Series) -> str:
    parts = [
        str(row.get("Date", "") or "").strip(),
        str(row.get("Amount", "") or "").strip(),
        str(row.get("Account Name", "") or "").strip(),
        str(row.get("Original Description", "") or "").strip()[:200],
    ]
    return "|".join(parts)


def assign_transaction_ids(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["Transaction ID"] = df.apply(
        lambda r: hashlib.sha256(transaction_fingerprint(r).encode("utf-8")).hexdigest()[:16],
        axis=1,
    )
    return df


def prepare_transaction_export_df(df: pd.DataFrame) -> pd.DataFrame:
    """Columns for transaction tabs: drop clutter, order Generated Description prominently."""
    out = df.drop(columns=["Amount_Numeric"], errors="ignore")
    out = out.drop(columns=[c for c in DROP_OUTPUT_COLUMNS if c in out.columns], errors="ignore")

    if "Generated Description" not in out.columns:
        out["Generated Description"] = out.apply(heuristic_generated_description, axis=1)

    # Place Generated Description after Date when possible
    preferred_front = [
        "Transaction ID",
        "Date",
        "Generated Description",
        "Merchant Key",
        "Amount",
        "Category",
        "Account Name",
        "Classification",
        "Expense Cadence",
        "In Monthly Run-Rate?",
        "Cadence Source",
    ]
    front = [c for c in preferred_front if c in out.columns]
    rest = [c for c in out.columns if c not in front]
    out = out[front + rest]
    if "Amount" in out.columns:
        if "Amount_Numeric" in df.columns:
            out["Amount"] = pd.to_numeric(df["Amount_Numeric"], errors="coerce")
        else:
            out["Amount"] = out["Amount"].apply(parse_amount)
    if "Date" in out.columns:
        if "Transaction Date" in df.columns:
            out["Date"] = pd.to_datetime(df["Transaction Date"], errors="coerce")
        else:
            out["Date"] = parse_transaction_dates(out["Date"])
    return out


def _hide_excel_columns(worksheet, column_names: tuple[str, ...], header_row: int = 1) -> None:
    from openpyxl.utils import get_column_letter

    headers = [cell.value for cell in worksheet[header_row]]
    for col_name in column_names:
        if col_name not in headers:
            continue
        col_idx = headers.index(col_name) + 1
        worksheet.column_dimensions[get_column_letter(col_idx)].hidden = True


def _coerce_cell_currency_value(value: object) -> float | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return parse_amount(str(value))
    except (TypeError, ValueError):
        return None


def _apply_currency_columns(
    worksheet,
    column_names: tuple[str, ...],
    *,
    header_row: int = 1,
) -> None:
    """Format named columns as US dollar amounts."""
    from openpyxl.utils import get_column_letter

    headers = [cell.value for cell in worksheet[header_row]]
    for col_name in column_names:
        if col_name not in headers:
            continue
        col_idx = headers.index(col_name) + 1
        letter = get_column_letter(col_idx)
        for row in range(header_row + 1, worksheet.max_row + 1):
            cell = worksheet[f"{letter}{row}"]
            amount = _coerce_cell_currency_value(cell.value)
            if amount is None:
                continue
            cell.value = amount
            cell.number_format = EXCEL_CURRENCY_FORMAT


def _percent_column_index(headers: list[Any], col_name: str) -> int | None:
    """Match header text allowing minor spacing variants (e.g. 'vs Avg%')."""
    target = str(col_name).strip().lower().replace(" ", "")
    for idx, header in enumerate(headers):
        key = str(header or "").strip().lower().replace(" ", "")
        if key == target:
            return idx + 1
    return None


def _apply_percent_columns(
    worksheet,
    column_names: tuple[str, ...],
    *,
    header_row: int = 1,
    last_row: int | None = None,
) -> None:
    """
    Format percent columns for Excel display.

    DataFrame values are whole numbers (42.3 = 42.3%); converted to fractions for Excel.
    """
    from openpyxl.utils import get_column_letter

    headers = [cell.value for cell in worksheet[header_row]]
    end_row = last_row if last_row is not None else int(worksheet.max_row or header_row)

    for col_name in column_names:
        col_idx = _percent_column_index(headers, col_name)
        if col_idx is None:
            continue
        letter = get_column_letter(col_idx)
        for row in range(header_row + 1, end_row + 1):
            cell = worksheet[f"{letter}{row}"]
            if cell.value is None or cell.value == "":
                continue
            try:
                val = float(cell.value)
            except (TypeError, ValueError):
                continue
            if abs(val) > 1.0:
                val = val / 100.0
            cell.value = val
            cell.number_format = EXCEL_PERCENT_FORMAT


def _coerce_cell_date_value(value: object) -> datetime | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    parsed = parse_transaction_dates(pd.Series([value])).iloc[0]
    if pd.isna(parsed):
        return None
    return parsed.to_pydatetime()


def _apply_short_date_columns(
    worksheet,
    column_names: tuple[str, ...],
    *,
    header_row: int = 1,
) -> None:
    """Format named columns as Excel short date (m/d/yyyy)."""
    from openpyxl.utils import get_column_letter

    headers = [cell.value for cell in worksheet[header_row]]
    for col_name in column_names:
        if col_name not in headers:
            continue
        col_idx = headers.index(col_name) + 1
        letter = get_column_letter(col_idx)
        for row in range(header_row + 1, worksheet.max_row + 1):
            cell = worksheet[f"{letter}{row}"]
            dt = _coerce_cell_date_value(cell.value)
            if dt is None:
                continue
            cell.value = dt
            cell.number_format = EXCEL_SHORT_DATE_FORMAT


def _display_cell_length(value: object, *, number_format: str | None = None) -> int:
    if value is None:
        return 0
    if isinstance(value, float) and pd.isna(value):
        return 0
    if number_format and number_format.startswith("$") and isinstance(value, (int, float)):
        return len(f"${abs(float(value)):,.2f}") + (1 if float(value) < 0 else 0)
    return len(str(value))


def _autofit_worksheet_columns(
    worksheet,
    *,
    header_row: int = 1,
    min_width: float = EXCEL_AUTOFIT_MIN_WIDTH,
    max_width: float = EXCEL_AUTOFIT_MAX_WIDTH,
    padding: float = EXCEL_AUTOFIT_PADDING,
) -> None:
    """Approximate Excel 'Autofit Column Width' for visible columns."""
    from openpyxl.utils import get_column_letter

    if not worksheet.max_column:
        return
    last_row = max(int(worksheet.max_row or 0), header_row)
    for col_idx in range(1, worksheet.max_column + 1):
        letter = get_column_letter(col_idx)
        dim = worksheet.column_dimensions[letter]
        if dim.hidden:
            continue
        max_len = 0
        for row in range(header_row, last_row + 1):
            cell = worksheet.cell(row, col_idx)
            max_len = max(
                max_len,
                _display_cell_length(cell.value, number_format=cell.number_format),
            )
        dim.width = min(max(max_len + padding, min_width), max_width)


def autofit_workbook(writer: pd.ExcelWriter) -> None:
    """Autofit column widths on every sheet in the workbook."""
    if not EXCEL_AUTOFIT_COLUMNS:
        return
    for worksheet in writer.sheets.values():
        _autofit_worksheet_columns(worksheet)


@contextmanager
def open_excel_workbook(path: Path | str) -> Iterator[pd.ExcelWriter]:
    """
    Context manager for Excel output. All sheets get column autofit on save (default).
    """
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        yield writer
        autofit_workbook(writer)


def sort_expenses_for_export(expense_df: pd.DataFrame) -> pd.DataFrame:
    """Sort expense rows for the Expenses tab: AI Category, then Generated Description."""
    if expense_df.empty:
        return expense_df
    sort_cols = [c for c in ("AI Category", "Generated Description") if c in expense_df.columns]
    if not sort_cols:
        return expense_df
    return expense_df.sort_values(sort_cols, na_position="last", kind="stable").reset_index(
        drop=True
    )


def _enable_worksheet_autofilter(worksheet) -> None:
    """Turn on Excel column filters on row 1 (Filter button in header)."""
    from openpyxl.utils import get_column_letter

    max_row = int(worksheet.max_row or 0)
    max_col = int(worksheet.max_column or 0)
    if max_row < 1 or max_col < 1:
        return
    worksheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"


def format_transaction_worksheet(
    worksheet,
    *,
    hide_descriptions: bool = False,
    short_date_columns: tuple[str, ...] = (),
    enable_autofilter: bool = False,
) -> None:
    if hide_descriptions:
        _hide_excel_columns(worksheet, HIDDEN_DESCRIPTION_COLUMNS)
    if short_date_columns:
        _apply_short_date_columns(worksheet, short_date_columns)
    _apply_currency_columns(worksheet, ("Amount",))
    worksheet.freeze_panes = "A2"
    if enable_autofilter:
        _enable_worksheet_autofilter(worksheet)


def format_baseline_worksheet(
    worksheet,
    *,
    overview_rows: int,
    categories_header_row: int,
    monthly_header_row: int | None = None,
) -> None:
    """Apply currency formats to the three Baseline tables."""
    headers = [cell.value for cell in worksheet[1]]
    if "Metric" in headers and "Value" in headers:
        metric_col = headers.index("Metric") + 1
        value_col = headers.index("Value") + 1
        for row in range(2, overview_rows + 2):
            metric = worksheet.cell(row, metric_col).value
            if metric not in BASELINE_OVERVIEW_CURRENCY_METRICS:
                continue
            cell = worksheet.cell(row, value_col)
            amount = _coerce_cell_currency_value(cell.value)
            if amount is not None:
                cell.value = amount
                cell.number_format = EXCEL_CURRENCY_FORMAT

    _apply_currency_columns(
        worksheet, BASELINE_CATEGORY_CURRENCY_COLUMNS, header_row=categories_header_row
    )
    if monthly_header_row is not None:
        _apply_currency_columns(
            worksheet, BASELINE_MONTHLY_CURRENCY_COLUMNS, header_row=monthly_header_row
        )


def format_spending_insights_worksheet(
    worksheet,
    *,
    overview_rows: int,
    overview_header_row: int,
    detail_header_rows: list[int],
    month_anchors: dict[str, int],
) -> None:
    """Currency formats and 'View categories' hyperlinks on the Summary insights sheet."""
    from openpyxl.styles import Font

    overview_last_row = overview_header_row + overview_rows
    _apply_currency_columns(
        worksheet,
        SUMMARY_OVERVIEW_CURRENCY_COLUMNS,
        header_row=overview_header_row,
    )
    _apply_percent_columns(
        worksheet,
        SUMMARY_OVERVIEW_PERCENT_COLUMNS,
        header_row=overview_header_row,
        last_row=overview_last_row,
    )

    sorted_detail_headers = sorted(detail_header_rows)
    for i, header_row in enumerate(sorted_detail_headers):
        if i + 1 < len(sorted_detail_headers):
            block_end = sorted_detail_headers[i + 1] - 2
        else:
            block_end = int(worksheet.max_row or header_row)
        _apply_currency_columns(
            worksheet,
            SUMMARY_DETAIL_CURRENCY_COLUMNS,
            header_row=header_row,
        )
        _apply_percent_columns(
            worksheet,
            SUMMARY_DETAIL_PERCENT_COLUMNS,
            header_row=header_row,
            last_row=block_end,
        )

    headers = [cell.value for cell in worksheet[overview_header_row]]
    if "View Categories" not in headers:
        return
    link_col = headers.index("View Categories") + 1
    month_col = headers.index("Budget Month") + 1 if "Budget Month" in headers else None
    link_font = Font(color="0563C1", underline="single")

    for row in range(overview_header_row + 1, overview_header_row + overview_rows + 1):
        month_val = (
            worksheet.cell(row, month_col).value if month_col is not None else None
        )
        month_key = str(month_val or "").strip()
        anchor = month_anchors.get(month_key)
        if not anchor:
            continue
        cell = worksheet.cell(row, link_col)
        cell.value = "View categories"
        cell.hyperlink = f"#'{SUMMARY_INSIGHTS_SHEET}'!A{anchor}"
        cell.font = link_font


def business_category_rule_keys(lookups: dict[str, pd.DataFrame]) -> set[str]:
    """Generated Description keys on BusinessCategoryRules (case-insensitive match)."""
    frame = lookups.get("BusinessCategoryRules")
    if frame is None or frame.empty or "Generated Description" not in frame.columns:
        return set()
    return {
        str(val).strip().lower()
        for val in frame["Generated Description"].fillna("").astype(str)
        if str(val).strip()
    }


def print_business_tagging_summary(df: pd.DataFrame) -> None:
    """Clarify business tagging vs bank Category column (common source of confusion)."""
    if "Classification" not in df.columns:
        return
    biz = df[df["Classification"].astype(str).str.strip().str.lower() == "business"]
    if biz.empty:
        return
    n_biz = len(biz)
    n_ai_be = 0
    if "AI Category" in biz.columns:
        n_ai_be = int((biz["AI Category"].astype(str).str.strip() == "Business Expenses").sum())
    n_bank_be = 0
    if "Category" in biz.columns:
        n_bank_be = int((biz["Category"].astype(str).str.strip() == "Business Expenses").sum())
    print(
        f"  Business tagging: {n_biz} row(s) with Classification=Business "
        f"(use this column to filter, not bank Category)",
        flush=True,
    )
    if "AI Category" in biz.columns:
        print(f"    AI Category 'Business Expenses': {n_ai_be}", flush=True)
    if "Category" in biz.columns and n_bank_be != n_biz:
        print(
            f"    Bank Category 'Business Expenses' only: {n_bank_be} "
            f"(e.g. Google/AWS labels; Zoom/Cursor use other bank categories)",
            flush=True,
        )


def mark_business_from_category_rules(
    df: pd.DataFrame,
    lookups: dict[str, pd.DataFrame],
) -> int:
    """
    Set Classification=Business when Generated Description matches BusinessCategoryRules.
    Bank exports often leave Classification as Personal; this applies your lookup list.
    """
    keys = business_category_rule_keys(lookups)
    if not keys or "Generated Description" not in df.columns:
        return 0
    if "Classification" not in df.columns:
        df["Classification"] = "Personal"

    marked = 0
    for idx, row in df.iterrows():
        desc = str(row.get("Generated Description", "") or "").strip().lower()
        if desc not in keys:
            continue
        if str(row.get("Classification", "") or "").strip().lower() != "business":
            df.at[idx, "Classification"] = "Business"
            marked += 1
    return marked


def _business_rules_lookup_keys(lookups: dict[str, pd.DataFrame]) -> set[str]:
    keys: set[str] = set()
    keys |= business_category_rule_keys(lookups)
    for mk in build_merchant_category_map(lookups):
        if mk:
            keys.add(mk)
    return keys


def business_rows_needing_rules(
    df: pd.DataFrame,
    lookups: dict[str, pd.DataFrame],
) -> pd.Index:
    known = _business_rules_lookup_keys(lookups)
    indices: list[int] = []
    for idx, row in df.iterrows():
        if not is_business_row(row):
            continue
        desc = str(row.get("Generated Description", "") or "").strip().lower()
        mk = _row_merchant_key(row).lower()
        if desc in known or mk in known:
            continue
        indices.append(int(idx))
    return pd.Index(indices)


def suggest_business_rules_batch(
    client: OpenAI,
    transactions: list[dict[str, Any]],
    model: str,
    *,
    use_json_mode: bool,
) -> list[dict[str, Any]]:
    user_content = json_for_prompt(transactions)
    messages = [
        {"role": "system", "content": BUSINESS_RULE_PROMPT},
        {
            "role": "user",
            "content": (
                "Suggest categorization rules for these business transactions. "
                'JSON {"results": [...]} only.\n\n'
                f"Transactions:\n{user_content}"
            ),
        },
    ]
    kwargs: dict[str, Any] = {
        "model": model,
        "temperature": 0.1,
        "messages": messages,
    }
    if use_json_mode:
        kwargs["response_format"] = {"type": "json_object"}

    response = client.chat.completions.create(**kwargs)
    raw = response.choices[0].message.content or "{}"
    parsed = extract_json_payload(raw)
    results = parsed.get("results", parsed if isinstance(parsed, list) else [])
    if not isinstance(results, list):
        raise ValueError(f"Unexpected business rule response: {raw[:500]}")
    return results


def enrich_business_lookup_rules(
    df: pd.DataFrame,
    client: OpenAI,
    model: str,
    lookups: dict[str, pd.DataFrame],
    *,
    batch_size: int,
    use_json_mode: bool,
) -> pd.DataFrame:
    """Return new BusinessCategoryRules rows for Business transactions missing lookup coverage."""
    need_idx = business_rows_needing_rules(df, lookups)
    if need_idx.empty:
        return pd.DataFrame(
            columns=[
                "Generated Description",
                "Source Category",
                "AI Category",
                "AI Sub-Category",
                "Budget Tier",
                "Type",
                "Sub-Type",
                "Classification",
                "Notes",
            ]
        )

    print(f"  Business lookup rules to analyze: {len(need_idx)}", flush=True)
    new_rows: list[dict[str, Any]] = []

    indices = need_idx.tolist()
    for start in range(0, len(indices), batch_size):
        batch_idx = indices[start : start + batch_size]
        payloads = []
        for i in batch_idx:
            row = df.loc[i]
            payloads.append(
                {
                    "index": int(i),
                    "generated_description": str(row.get("Generated Description", ""))[:200],
                    "original_description": str(row.get("Original Description", ""))[:300],
                    "category": str(row.get("Category", "")),
                    "amount": str(row.get("Amount", "")),
                    "account": str(row.get("Account Name", "")),
                }
            )
        try:
            results = suggest_business_rules_batch(
                client, payloads, model, use_json_mode=use_json_mode
            )
        except Exception as exc:
            print(f"  Business rule LLM failed; using row defaults: {exc}", flush=True)
            results = []

        result_by_idx = {int(r.get("index", -1)): r for r in results}
        for i in batch_idx:
            row = df.loc[i]
            r = result_by_idx.get(int(i), {})
            desc = str(row.get("Generated Description", "") or "").strip()
            new_rows.append(
                {
                    "Generated Description": desc,
                    "Source Category": str(row.get("Category", "") or ""),
                    "AI Category": str(
                        r.get("ai_category", row.get("AI Category", "")) or ""
                    ).strip(),
                    "AI Sub-Category": str(
                        r.get("ai_sub_category", row.get("AI Sub-Category", desc)) or ""
                    ).strip(),
                    "Budget Tier": str(
                        r.get("budget_tier", row.get("Budget Tier", "")) or ""
                    ).strip(),
                    "Type": str(r.get("type", row.get("Type", "Variable")) or "Variable").strip(),
                    "Sub-Type": str(row.get("Sub-Type", "") or ""),
                    "Classification": "Business",
                    "Notes": str(r.get("notes", "Auto-added from Business transaction") or ""),
                }
            )

    return pd.DataFrame(new_rows)


def apply_business_rules_to_df(
    df: pd.DataFrame,
    rules: pd.DataFrame,
    *,
    spend_mask: pd.Series | None = None,
) -> int:
    """Apply BusinessCategoryRules rows to matching Business transactions in-place."""
    if rules is None or rules.empty or "Generated Description" not in rules.columns:
        return 0
    by_desc: dict[str, pd.Series] = {}
    for _, brow in rules.iterrows():
        key = str(brow.get("Generated Description", "") or "").strip().lower()
        if key:
            by_desc[key] = brow
    updated = 0
    for idx, row in df.iterrows():
        if not is_business_row(row):
            continue
        key = str(row.get("Generated Description", "") or "").strip().lower()
        match = by_desc.get(key)
        if match is None:
            continue
        for col, target in (
            ("AI Category", "AI Category"),
            ("AI Sub-Category", "AI Sub-Category"),
            ("Type", "Type"),
            ("Sub-Type", "Sub-Type"),
            ("Flow Type", "Flow Type"),
        ):
            val = match.get(col, "")
            if pd.notna(val) and str(val).strip() not in ("", "nan"):
                df.at[idx, target] = str(val).strip()
        if spend_mask is None or bool(spend_mask.loc[idx]):
            tier = match.get("Budget Tier", "")
            if pd.notna(tier) and str(tier).strip() not in ("", "nan", "Review"):
                df.at[idx, "Budget Tier"] = str(tier).strip()
        updated += 1
    return updated


def merge_business_category_rules(
    existing: pd.DataFrame | None,
    new_rules: pd.DataFrame,
) -> pd.DataFrame:
    cols = [
        "Generated Description",
        "Source Category",
        "AI Category",
        "AI Sub-Category",
        "Budget Tier",
        "Type",
        "Sub-Type",
        "Classification",
        "Notes",
    ]
    if new_rules is None or new_rules.empty:
        if existing is None or existing.empty:
            return pd.DataFrame(columns=cols)
        return existing

    if existing is None or existing.empty:
        out = new_rules.copy()
    else:
        existing = existing.copy()
        for col in cols:
            if col not in existing.columns:
                existing[col] = ""
        new_rules = new_rules.copy()
        known = set(existing["Generated Description"].fillna("").astype(str).str.strip().str.lower())
        append = new_rules[
            ~new_rules["Generated Description"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.lower()
            .isin(known)
        ]
        out = pd.concat([existing[cols], append[cols]], ignore_index=True)

    out["Generated Description"] = out["Generated Description"].fillna("").astype(str)
    out = out.drop_duplicates(subset=["Generated Description"], keep="first")
    return out.sort_values("Generated Description").reset_index(drop=True)


def empty_custom_rules_sheet() -> pd.DataFrame:
    # object columns: Excel empty cells otherwise become float64 and reject JSON strings
    return pd.DataFrame({col: pd.Series(dtype=object) for col in CUSTOM_RULES_COLUMNS})


def _custom_rules_cell_str(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def normalize_custom_rules_sheet(frame: pd.DataFrame | None) -> pd.DataFrame:
    if frame is None or frame.empty:
        return empty_custom_rules_sheet()
    out = frame.copy()
    if "Rule" not in out.columns:
        # Legacy or mistaken header row
        return empty_custom_rules_sheet()
    for col in CUSTOM_RULES_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    out = out[list(CUSTOM_RULES_COLUMNS)]
    for col in CUSTOM_RULES_COLUMNS:
        out[col] = out[col].map(_custom_rules_cell_str)
    return out


def compile_custom_rule_text(
    client: OpenAI,
    rule_text: str,
    model: str,
    *,
    use_json_mode: bool,
) -> tuple[dict[str, Any] | None, str]:
    """Ask LLM to convert freeform Rule column text into executable JSON."""
    messages = [
        {"role": "system", "content": CUSTOM_RULE_COMPILER_PROMPT},
        {
            "role": "user",
            "content": (
                "Convert this rule to JSON.\n\n"
                f"Rule:\n{rule_text.strip()}"
            ),
        },
    ]
    kwargs: dict[str, Any] = {
        "model": model,
        "temperature": 0.1,
        "messages": messages,
    }
    if use_json_mode:
        kwargs["response_format"] = {"type": "json_object"}

    response = client.chat.completions.create(**kwargs)
    raw = response.choices[0].message.content or "{}"
    try:
        parsed = extract_json_payload(raw)
    except json.JSONDecodeError as exc:
        return None, f"Invalid JSON from model: {exc}"

    rule = parsed.get("rule", parsed) if isinstance(parsed, dict) else parsed
    if not isinstance(rule, dict):
        return None, f"Expected object with 'rule' key; got: {str(parsed)[:300]}"
    if "rule_type" not in rule:
        return None, "Compiled rule missing rule_type"
    return rule, ""


def compile_custom_rules_sheet(
    sheet: pd.DataFrame,
    client: OpenAI,
    model: str,
    *,
    use_json_mode: bool,
) -> pd.DataFrame:
    """Compile Pending rows via LLM; update Compiled Rule, Status, Last Error."""
    out = normalize_custom_rules_sheet(sheet)
    if out.empty:
        return out

    now = pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")
    for idx, row in out.iterrows():
        status = str(row.get("Status", "") or "").strip().lower()
        rule_text = str(row.get("Rule", "") or "").strip()
        if not rule_text:
            continue
        if status == CUSTOM_RULE_STATUS_DISABLED.lower():
            continue
        compiled_existing = str(row.get("Compiled Rule", "") or "").strip()
        if status == CUSTOM_RULE_STATUS_ACTIVE.lower() and compiled_existing:
            continue
        if status not in (
            "",
            CUSTOM_RULE_STATUS_PENDING.lower(),
            "compile",
            CUSTOM_RULE_STATUS_ERROR.lower(),
        ):
            continue

        print(f"  Compiling custom rule row {idx + 1}...", flush=True)
        try:
            compiled, err = compile_custom_rule_text(
                client, rule_text, model, use_json_mode=use_json_mode
            )
        except Exception as exc:
            compiled, err = None, str(exc)

        if compiled is not None:
            out.at[idx, "Compiled Rule"] = json.dumps(compiled, separators=(",", ":"))
            out.at[idx, "Status"] = CUSTOM_RULE_STATUS_ACTIVE
            out.at[idx, "Last Error"] = ""
        else:
            out.at[idx, "Status"] = CUSTOM_RULE_STATUS_ERROR
            out.at[idx, "Last Error"] = err[:500]
        out.at[idx, "Updated At"] = now

    return out


def load_active_custom_rules(sheet: pd.DataFrame | None) -> list[dict[str, Any]]:
    rules: list[dict[str, Any]] = []
    out = normalize_custom_rules_sheet(sheet)
    for _, row in out.iterrows():
        if str(row.get("Status", "") or "").strip().lower() != CUSTOM_RULE_STATUS_ACTIVE.lower():
            continue
        raw = str(row.get("Compiled Rule", "") or "").strip()
        if not raw:
            continue
        try:
            rule = json.loads(raw)
        except json.JSONDecodeError:
            continue
        if isinstance(rule, dict) and rule.get("rule_type"):
            rules.append(rule)
    return rules


def _apply_custom_field_spec(df: pd.DataFrame, idx: int, spec: dict[str, Any]) -> None:
    for key, val in spec.items():
        if val is None or (isinstance(val, float) and pd.isna(val)):
            continue
        col = CUSTOM_RULE_FIELD_MAP.get(str(key).strip().lower(), str(key))
        if col not in df.columns:
            continue
        text = str(val).strip()
        if text and text.lower() != "nan":
            df.at[idx, col] = text


def _normalize_match_patterns(value: Any) -> list[str]:
    if value is None:
        return []
    items = value if isinstance(value, (list, tuple)) else [value]
    patterns: list[str] = []
    for item in items:
        text = str(item or "").strip()
        if text and text.lower() != "nan":
            patterns.append(text)
    return patterns


def _match_text_pattern(series: pd.Series, pattern: str) -> pd.Series:
    """Case-insensitive match of one pattern against a lowercase text series."""
    target = str(pattern or "").strip().lower()
    if not target:
        return pd.Series(False, index=series.index)
    if target.startswith("*") and target.endswith("*") and len(target) > 2:
        needle = target[1:-1]
        return series.str.contains(needle, regex=False) if needle else pd.Series(False, index=series.index)
    if target.endswith("*") and len(target) > 1:
        return series.str.startswith(target[:-1])
    if target.startswith("*") and len(target) > 1:
        return series.str.endswith(target[1:])
    return series == target


def _match_patterns_on_series(series: pd.Series, patterns: Any) -> pd.Series:
    normalized = _normalize_match_patterns(patterns)
    if not normalized:
        return pd.Series(False, index=series.index)
    mask = pd.Series(False, index=series.index)
    lowered = series.fillna("").astype(str).str.strip().str.lower()
    for pattern in normalized:
        mask |= _match_text_pattern(lowered, pattern)
    return mask


def _match_generated_description(df: pd.DataFrame, desc: Any) -> pd.Series:
    if "Generated Description" not in df.columns:
        return pd.Series(False, index=df.index)
    series = df["Generated Description"]
    return _match_patterns_on_series(series, desc)


def _match_any_description(df: pd.DataFrame, desc: Any) -> pd.Series:
    columns = (
        "Generated Description",
        "Original Description",
        "Simple Description",
        "User Description",
    )
    present = [col for col in columns if col in df.columns]
    if not present:
        return pd.Series(False, index=df.index)
    mask = pd.Series(False, index=df.index)
    for col in present:
        mask |= _match_patterns_on_series(df[col], desc)
    return mask


def _match_custom_rule_amount(df: pd.DataFrame, amount_spec: Any) -> pd.Series:
    """Match rows where |Amount| equals the target (expenses are often negative)."""
    if amount_spec is None or (isinstance(amount_spec, float) and pd.isna(amount_spec)):
        return pd.Series(True, index=df.index)
    text = str(amount_spec).strip().replace("$", "").replace(",", "")
    if not text or text.lower() == "nan":
        return pd.Series(True, index=df.index)
    try:
        target = abs(parse_amount(text))
    except (TypeError, ValueError):
        return pd.Series(False, index=df.index)
    if "Amount_Numeric" not in df.columns:
        amounts = df.get("Amount", pd.Series("", index=df.index)).apply(parse_amount)
    else:
        amounts = df["Amount_Numeric"]
    return amounts.apply(lambda a: abs(float(a)) == target)


def _custom_rule_match_mask(df: pd.DataFrame, match: dict[str, Any]) -> pd.Series:
    """AND of all match keys in a compiled custom rule."""
    match = match or {}
    mask = pd.Series(True, index=df.index)
    if "description" in match:
        mask &= _match_any_description(df, match.get("description"))
    if "generated_description" in match:
        mask &= _match_generated_description(df, match.get("generated_description"))
    if "amount" in match:
        mask &= _match_custom_rule_amount(df, match.get("amount"))
    return mask


def _apply_custom_rule_assign(df: pd.DataFrame, rule: dict[str, Any]) -> int:
    match = rule.get("match", {}) or {}
    mask = _custom_rule_match_mask(df, match)
    spec = rule.get("set", {}) or {}
    if not mask.any() or not spec:
        return 0
    for idx in df[mask].index:
        _apply_custom_field_spec(df, int(idx), spec)
    return int(mask.sum())


def _apply_custom_rule_monthly_split_max(df: pd.DataFrame, rule: dict[str, Any]) -> int:
    match = rule.get("match", {}) or {}
    mask = _custom_rule_match_mask(df, match)
    if not mask.any():
        return 0

    group_col = str(rule.get("group_by", "Budget Month") or "Budget Month")
    if group_col not in df.columns:
        group_col = "Calendar Month"
    min_rows = int(rule.get("min_rows_per_group", 2))
    when_max = rule.get("when_max", {}) or {}
    when_other = rule.get("when_other", {}) or {}

    updated = 0
    subset = df[mask]
    for _, grp in subset.groupby(group_col, dropna=False):
        if len(grp) < min_rows:
            continue
        amounts = grp["Amount_Numeric"].abs() if "Amount_Numeric" in grp.columns else grp.index.to_series()
        idx_max = amounts.idxmax()
        for idx in grp.index:
            spec = when_max if idx == idx_max else when_other
            if spec:
                _apply_custom_field_spec(df, int(idx), spec)
                updated += 1
    return updated


def apply_custom_rules(df: pd.DataFrame, compiled_rules: list[dict[str, Any]]) -> int:
    """
    Apply Active compiled custom rules.

    Called last in the pipeline so CustomRules override CategoryRules, Categories,
    BusinessCategoryRules, LLM review, and expense cadence lookup on shared fields.
    """
    total = 0
    for rule in compiled_rules:
        rule_type = str(rule.get("rule_type", "") or "").strip().lower()
        if rule_type == "assign":
            total += _apply_custom_rule_assign(df, rule)
        elif rule_type in ("monthly_split_max", "monthly_highest_split"):
            total += _apply_custom_rule_monthly_split_max(df, rule)
    return total


def empty_expense_cadence_rules_sheet() -> pd.DataFrame:
    return pd.DataFrame(columns=list(EXPENSE_CADENCE_RULE_COLUMNS))


def normalize_expense_cadence_rules_sheet(frame: pd.DataFrame | None) -> pd.DataFrame:
    if frame is None or frame.empty:
        return empty_expense_cadence_rules_sheet()
    out = frame.copy()
    if "Generated Description" not in out.columns:
        return empty_expense_cadence_rules_sheet()
    for col in EXPENSE_CADENCE_RULE_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    return out[list(EXPENSE_CADENCE_RULE_COLUMNS)]


def merge_expense_cadence_rules(
    existing: pd.DataFrame | None,
    new_rules: pd.DataFrame,
) -> pd.DataFrame:
    """Preserve existing ExpenseCadenceRules; append new Generated Description keys only."""
    cols = list(EXPENSE_CADENCE_RULE_COLUMNS)
    if new_rules is None or new_rules.empty:
        return normalize_expense_cadence_rules_sheet(existing)

    new_rules = normalize_expense_cadence_rules_sheet(new_rules)
    if existing is None or existing.empty:
        return new_rules

    old = normalize_expense_cadence_rules_sheet(existing)
    old_keys = {
        str(r.get("Generated Description", "") or "").strip().lower()
        for _, r in old.iterrows()
        if str(r.get("Generated Description", "") or "").strip()
    }
    append_rows = []
    for _, row in new_rules.iterrows():
        key = str(row.get("Generated Description", "") or "").strip().lower()
        if key and key not in old_keys:
            append_rows.append(row)
    if not append_rows:
        return old
    return pd.concat([old, pd.DataFrame(append_rows)], ignore_index=True)


def _normalize_cadence_label(value: Any) -> str:
    text = str(value or "").strip()
    if not text or text.lower() == "nan":
        return CADENCE_UNKNOWN
    lower = text.lower()
    if lower in ("monthly", "month"):
        return CADENCE_MONTHLY
    if lower in ("yearly", "annual", "year"):
        return CADENCE_YEARLY
    if lower in ("one-time", "one time", "onetime"):
        return CADENCE_ONETIME
    if lower in ("unplanned", "unexpected"):
        return CADENCE_UNPLANNED
    if text in CADENCE_VALID:
        return text
    return CADENCE_UNKNOWN


def _normalize_runrate_flag(value: Any, cadence: str) -> str:
    text = str(value or "").strip().upper()
    if text in ("Y", "YES", "TRUE", "1"):
        return "Y"
    if text in ("N", "NO", "FALSE", "0"):
        return "N"
    return CADENCE_DEFAULT_RUNRATE.get(cadence, "Y")


def init_expense_cadence_columns(df: pd.DataFrame) -> None:
    df["Expense Cadence"] = CADENCE_UNKNOWN
    df["In Monthly Run-Rate?"] = ""
    df["Cadence Source"] = ""
    df["Cadence Note"] = ""


def apply_expense_cadence_lookup(
    df: pd.DataFrame,
    rules: pd.DataFrame | None,
    *,
    spend_mask: pd.Series,
) -> int:
    """Apply ExpenseCadenceRules by Generated Description (case-insensitive)."""
    rules = normalize_expense_cadence_rules_sheet(rules)
    if rules.empty or "Generated Description" not in df.columns:
        return 0

    by_desc: dict[str, pd.Series] = {}
    for _, row in rules.iterrows():
        key = str(row.get("Generated Description", "") or "").strip().lower()
        if key:
            by_desc[key] = row

    updated = 0
    for idx in df[spend_mask].index:
        key = str(df.at[idx, "Generated Description"] or "").strip().lower()
        match = by_desc.get(key)
        if match is None:
            continue
        cadence = _normalize_cadence_label(match.get("Cadence", ""))
        runrate = _normalize_runrate_flag(match.get("In Monthly Run-Rate?", ""), cadence)
        df.at[idx, "Expense Cadence"] = cadence
        df.at[idx, "In Monthly Run-Rate?"] = runrate
        df.at[idx, "Cadence Source"] = CADENCE_SOURCE_LOOKUP
        note = str(match.get("Notes", "") or "").strip()
        if note:
            df.at[idx, "Cadence Note"] = note
        updated += 1
    return updated


def build_analytics_ledger(
    current_df: pd.DataFrame,
    history_path: Path | None,
) -> pd.DataFrame:
    """Merge cumulative history with the current run for cadence pattern detection."""
    parts: list[pd.DataFrame] = []
    if history_path is not None and history_path.exists():
        hist_sheets = load_history_workbook(history_path)
        prior = rebuild_ledger_from_history(hist_sheets)
        if not prior.empty:
            parts.append(prior)

    cur = current_df.copy()
    if "Amount_Numeric" not in cur.columns and "Amount" in cur.columns:
        cur["Amount_Numeric"] = cur["Amount"].apply(parse_amount)
    parts.append(cur)

    if not parts:
        return pd.DataFrame()
    if len(parts) == 1:
        ledger = parts[0]
    else:
        ledger = pd.concat(parts, ignore_index=True)
        if "Transaction ID" in ledger.columns:
            ledger = ledger.drop_duplicates(subset=["Transaction ID"], keep="last")

    if "Budget Month" not in ledger.columns and "Transaction Date" in ledger.columns:
        ledger["Budget Month"] = pd.to_datetime(
            ledger["Transaction Date"], errors="coerce"
        ).dt.to_period("M").astype(str)
    return ledger


def analyze_merchant_cadence_profiles(ledger: pd.DataFrame) -> dict[str, dict[str, Any]]:
    """
    Per Generated Description, infer monthly vs yearly vs one-time vs unplanned
    from multi-month spend totals in the analytics ledger.
    """
    spend = _spend_rows(ledger)
    if spend.empty or "Generated Description" not in spend.columns:
        return {}

    monthly = (
        spend.groupby(["Generated Description", "Budget Month"], dropna=False)["Spend Amount"]
        .sum()
        .reset_index()
    )

    profiles: dict[str, dict[str, Any]] = {}
    for desc, grp in monthly.groupby("Generated Description", dropna=False):
        desc_key = str(desc).strip().lower()
        if not desc_key:
            continue
        by_month = grp.set_index("Budget Month")["Spend Amount"]
        if len(by_month) < CADENCE_MIN_HISTORY_MONTHS:
            continue

        typical = float(by_month.median())
        if typical <= 0:
            continue

        threshold = typical * CADENCE_SPIKE_RATIO
        high_months = [str(m) for m, total in by_month.items() if float(total) >= threshold]

        pattern = "monthly"
        note = f"Typical month ~${typical:,.2f}"
        if len(high_months) >= 2:
            try:
                periods = sorted(pd.Period(m, freq="M") for m in high_months)
                gaps = [(periods[i + 1] - periods[i]).n for i in range(len(periods) - 1)]
            except Exception:
                gaps = []
            if any(CADENCE_YEARLY_GAP_MIN <= g <= CADENCE_YEARLY_GAP_MAX for g in gaps):
                pattern = "yearly"
                note = (
                    f"{len(high_months)} spike month(s); ~annual pattern "
                    f"(typical month ~${typical:,.2f})"
                )
            else:
                pattern = "unplanned"
                note = (
                    f"{len(high_months)} spike month(s); no annual spacing "
                    f"(typical ~${typical:,.2f})"
                )
        elif len(high_months) == 1:
            pattern = "onetime"
            note = f"Single spike month in history (typical ~${typical:,.2f})"

        profiles[desc_key] = {
            "pattern": pattern,
            "typical_monthly": typical,
            "high_months": set(high_months),
            "note": note,
        }
    return profiles


def _cadence_from_pattern(pattern: str) -> str:
    if pattern == "yearly":
        return CADENCE_YEARLY
    if pattern == "onetime":
        return CADENCE_ONETIME
    if pattern == "unplanned":
        return CADENCE_UNPLANNED
    return CADENCE_MONTHLY


def apply_detected_expense_cadence(
    df: pd.DataFrame,
    profiles: dict[str, dict[str, Any]],
    *,
    spend_mask: pd.Series,
) -> int:
    """Tag spend rows from history-based detection (skips Lookup-tagged rows)."""
    if not profiles:
        return 0

    updated = 0
    for idx in df[spend_mask].index:
        if str(df.at[idx, "Cadence Source"] or "").strip() == CADENCE_SOURCE_LOOKUP:
            continue

        desc_key = str(df.at[idx, "Generated Description"] or "").strip().lower()
        profile = profiles.get(desc_key)
        if profile is None:
            continue

        month = str(df.at[idx, "Budget Month"] or "")
        amount = abs(float(df.at[idx, "Amount_Numeric"]))
        typical = float(profile["typical_monthly"])
        threshold = typical * CADENCE_SPIKE_RATIO

        if month in profile["high_months"] and amount >= threshold * 0.5:
            cadence = _cadence_from_pattern(str(profile["pattern"]))
        else:
            cadence = CADENCE_MONTHLY

        df.at[idx, "Expense Cadence"] = cadence
        df.at[idx, "In Monthly Run-Rate?"] = CADENCE_DEFAULT_RUNRATE.get(cadence, "Y")
        df.at[idx, "Cadence Source"] = CADENCE_SOURCE_DETECTED
        df.at[idx, "Cadence Note"] = str(profile.get("note", "") or "")
        updated += 1
    return updated


def finalize_expense_cadence_defaults(df: pd.DataFrame, *, spend_mask: pd.Series) -> None:
    """Unknown cadence defaults to included in monthly run-rate until lookup/detection sets otherwise."""
    for idx in df[spend_mask].index:
        if str(df.at[idx, "Cadence Source"] or "").strip():
            continue
        df.at[idx, "Expense Cadence"] = CADENCE_UNKNOWN
        df.at[idx, "In Monthly Run-Rate?"] = "Y"
        df.at[idx, "Cadence Source"] = CADENCE_SOURCE_DEFAULT


def build_cadence_review_df(df: pd.DataFrame) -> pd.DataFrame:
    """Rows auto-tagged as irregular (Detected, not Monthly) for user review."""
    if "Cadence Source" not in df.columns:
        return pd.DataFrame()
    mask = (
        (df["Cadence Source"].astype(str) == CADENCE_SOURCE_DETECTED)
        & (df["Expense Cadence"] != CADENCE_MONTHLY)
        & (df["Include in Spend?"] == "Y")
    )
    review = df.loc[mask].copy()
    if review.empty:
        return pd.DataFrame()

    cols = [
        c
        for c in (
            "Transaction ID",
            "Transaction Date",
            "Budget Month",
            "Generated Description",
            "Amount",
            "AI Category",
            "AI Sub-Category",
            "Expense Cadence",
            "In Monthly Run-Rate?",
            "Cadence Source",
            "Cadence Note",
        )
        if c in review.columns
    ]
    out = review[cols].copy()
    if "Transaction Date" in out.columns:
        out = out.sort_values(["Budget Month", "Generated Description", "Transaction Date"])
    return out.reset_index(drop=True)


def _abs_spend_sum(spend: pd.DataFrame, row_mask: pd.Series) -> float:
    if spend.empty or not row_mask.any():
        return 0.0
    return float(spend.loc[row_mask, "Amount_Numeric"].abs().sum())


def budget_tier_from_category(category: str) -> str:
    """Map category text to your 3-tier budget model (Need/Want/Wish)."""
    c = str(category or "").lower()

    need_patterns = [
        "mortgage",
        "hoa",
        "rent",
        "housing",
        "utilities",
        "telephone",
        "internet",
        "insurance",
        "health",
        "medical",
        "household repairs",
        "repairs",
        "dues and subscriptions",
        "dues",
        "services",
        "postage",
        "automotive expenses",
        "automotive",
    ]
    if any(p in c for p in need_patterns):
        return "Need"

    want_patterns = [
        "grocer",
        "gasoline",
        "fuel",
        "pets",
        "personal care",
        "gym",
        "clothing",
        "electronics",
        "education",
        "warranty",
    ]
    if any(p in c for p in want_patterns):
        return "Want"

    wish_patterns = [
        "dining",
        "restaurant",
        "coffee",
        "entertain",
        "travel",
        "hobbies",
        "gift",
        "charitable",
        "general merchandise",
        "shopping",
        "online services",
    ]
    if any(p in c for p in wish_patterns):
        return "Wish"

    return "Review"


def cost_type_from_category(category: str) -> str:
    """Fixed vs Variable default (LLM can override for Review items)."""
    c = str(category or "").lower()
    fixed_patterns = [
        "mortgage",
        "hoa",
        "rent",
        "utilities",
        "telephone",
        "internet",
        "insurance",
        "dues and subscriptions",
        "dues",
        "subscription",
        "online services",
    ]
    if any(p in c for p in fixed_patterns):
        return "Fixed"

    # Gym membership tends to be recurring
    if "gym" in c:
        return "Fixed"

    return "Variable"


def is_paycheck_row(row: pd.Series) -> bool:
    cat = str(row.get("Category", "") or "")
    amt = float(row.get("Amount_Numeric", 0.0))
    if cat == "Paychecks/Salary" and amt > 0:
        return True
    combined = " ".join(
        [str(row.get("Original Description", "") or ""), str(row.get("Simple Description", "") or "")]
    ).lower()
    return "payroll" in combined and amt > 0


def _description_suggests_refund(row: pd.Series) -> bool:
    """Heuristic: bank text looks like a return/refund/reversal."""
    text = " ".join(
        [
            str(row.get("Original Description", "") or ""),
            str(row.get("Simple Description", "") or ""),
            str(row.get("User Description", "") or ""),
            str(row.get("Generated Description", "") or ""),
        ]
    ).lower()
    hints = (
        "refund",
        "return",
        "reversal",
        "reversed",
        "chargeback",
        "credit adj",
        "purchase return",
        "merchandise return",
    )
    return any(h in text for h in hints)


def flow_type_from_row(row: pd.Series) -> str:
    """Classify how the amount should be summarized."""
    cat = str(row.get("Category", "") or "").strip()
    amt = float(row.get("Amount_Numeric", 0.0))

    internal = {"Transfers", "Credit Card Payments", "Savings", "Securities Trades"}
    true_income = {"Paychecks/Salary", "Interest"}
    adjustment = {"Refunds/Adjustments", "Rewards", "Other Income", "Expense Reimbursement", "Deposits"}

    if cat in internal:
        return "Transfer"

    if amt > 0:
        if cat in true_income:
            return "Income"
        if cat in adjustment:
            return "Adjustment"
        if _description_suggests_refund(row):
            return "Adjustment"
        # Positive amount on a spending category (e.g. JetBrains subscription refund)
        # is a credit against prior spend — not wages or business income.
        return "Adjustment"

    # amt < 0
    return "Expense"


def _workbook_path(filename: str) -> Path:
    """Resolve lookup/history workbooks; prefer scripts/, migrate legacy copies from project root."""
    path = Path(filename)
    if path.is_absolute():
        return path
    if len(path.parts) > 1:
        canonical = (PROJECT_ROOT / path).resolve()
    else:
        canonical = (SCRIPTS_DIR / path.name).resolve()

    legacy = (PROJECT_ROOT / path.name).resolve()
    if not canonical.exists() and legacy.exists() and legacy != canonical:
        try:
            canonical.parent.mkdir(parents=True, exist_ok=True)
            legacy.rename(canonical)
            print(f"  Relocated {legacy.name} → {canonical.parent.name}/", flush=True)
        except OSError:
            return legacy
    return canonical


def default_lookup_path() -> Path:
    """Shared lookup workbook in scripts/ (rules, description cache, custom rules)."""
    return _workbook_path(LOOKUP_FILENAME)


def default_history_path() -> Path:
    """
    Cumulative history in scripts/ — read for cadence detection when present;
    written only with --update-history (not a per-run export; see output/ for those).
    """
    return _workbook_path(HISTORY_FILENAME)


def annotate_history_run(export_df: pd.DataFrame, *, source_file: str) -> pd.DataFrame:
    """Tag rows with source CSV and processing time for the cumulative history file."""
    out = export_df.copy()
    out["Source File"] = source_file
    out["Processed At"] = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    return out


def merge_dedupe_transactions(
    existing: pd.DataFrame | None,
    new_rows: pd.DataFrame,
) -> pd.DataFrame:
    """
    Merge transaction rows; duplicate Transaction ID keeps the latest Processed At.
    Re-running the same CSV replaces prior rows instead of appending duplicates.
    """
    if new_rows is None or new_rows.empty:
        if existing is None or existing.empty:
            return pd.DataFrame()
        combined = existing.copy()
    elif existing is None or existing.empty:
        combined = new_rows.copy()
    else:
        combined = pd.concat([existing, new_rows], ignore_index=True)

    if combined.empty:
        return combined

    if "Transaction ID" not in combined.columns:
        raise ValueError("History merge requires a Transaction ID column")

    combined["Transaction ID"] = combined["Transaction ID"].astype(str)
    if "Processed At" in combined.columns:
        combined = combined.sort_values("Processed At", na_position="last")
    combined = combined.drop_duplicates(subset=["Transaction ID"], keep="last")

    sort_cols = [c for c in ("Date", "Generated Description") if c in combined.columns]
    if sort_cols:
        combined = combined.sort_values(sort_cols, na_position="last")
    return combined.reset_index(drop=True)


def load_history_workbook(history_path: Path) -> dict[str, pd.DataFrame]:
    sheets: dict[str, pd.DataFrame] = {}
    if not history_path.exists():
        return sheets
    try:
        xl = pd.ExcelFile(history_path)
    except Exception:
        return sheets
    for name in HISTORY_SHEETS:
        if name in xl.sheet_names:
            sheets[name] = pd.read_excel(xl, sheet_name=name)
    return sheets


def rebuild_ledger_from_history(sheets: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Rebuild an in-memory ledger (with Amount_Numeric) from cumulative transaction sheets."""
    parts: list[pd.DataFrame] = []
    for name in ("Income", "Expenses", "Adjustments"):
        frame = sheets.get(name)
        if frame is not None and not frame.empty:
            parts.append(frame.copy())
    if not parts:
        return pd.DataFrame()

    ledger = pd.concat(parts, ignore_index=True)
    if "Amount" in ledger.columns:
        ledger["Amount_Numeric"] = ledger["Amount"].apply(parse_amount)
    if "Transaction Date" not in ledger.columns and "Date" in ledger.columns:
        ledger["Transaction Date"] = parse_transaction_dates(ledger["Date"])
    return ledger


def _category_avg_monthly_spend(spend: pd.DataFrame) -> pd.Series:
    """
    Average monthly spend per AI Category (mean of each month's category total).

    Not per transaction — e.g. Restaurants/Dining is ~$2k/mo, not ~$39/check.
    """
    if spend.empty or "Budget Month" not in spend.columns:
        return pd.Series(dtype=float)
    monthly_by_cat = (
        spend.groupby(["AI Category", "Budget Month"], dropna=False)["Spend Amount"]
        .sum()
    )
    return monthly_by_cat.groupby("AI Category", dropna=False).mean()


def build_monthly_expense_overview(
    df: pd.DataFrame,
    spend: pd.DataFrame,
) -> pd.DataFrame:
    """
    Month totals vs average monthly total expenses.
    Averages use all budget months present in the dataset (this export or history file).
    """
    empty = pd.DataFrame(
        columns=[
            "Budget Month",
            "Gross Income",
            "Total Expenses",
            "Avg Monthly Expenses",
            "vs Avg",
            "vs Avg %",
            "View Categories",
        ]
    )
    if spend.empty or "Budget Month" not in spend.columns:
        return empty

    monthly_totals = spend.groupby("Budget Month", dropna=False)["Spend Amount"].sum()
    overall_avg = float(monthly_totals.mean()) if len(monthly_totals) else 0.0
    budget_months = sorted(monthly_totals.index.tolist())

    rows: list[dict[str, Any]] = []
    for m in budget_months:
        total = float(monthly_totals.get(m, 0.0))
        vs_avg = total - overall_avg
        vs_pct = (vs_avg / overall_avg * 100) if overall_avg else 0.0

        income_budget = 0.0
        if "Budget Month" in df.columns:
            sub = df[df["Budget Month"] == m]
            income_budget = float(
                sub[
                    (sub["Flow Type"] == "Income") & (sub["Amount_Numeric"] > 0)
                ]["Amount_Numeric"].sum()
            )

        rows.append(
            {
                "Budget Month": m,
                "Gross Income": round(income_budget, 2),
                "Total Expenses": round(total, 2),
                "Avg Monthly Expenses": round(overall_avg, 2),
                "vs Avg": round(vs_avg, 2),
                "vs Avg %": round(vs_pct, 1),
                "View Categories": "",
            }
        )

    return pd.DataFrame(rows)


def build_month_category_breakdown(
    spend: pd.DataFrame,
    budget_month: str,
    *,
    cat_avg: pd.Series,
    month_total: float,
) -> pd.DataFrame:
    """
    Categories for one month: sorted by deviation from category average (largest first),
    then by month spend descending.
    """
    cols = [
        "AI Category",
        "Month Spend",
        "Avg Monthly Spend",
        "vs Avg",
        "vs Avg %",
        "% of Month",
    ]
    month_spend = spend[spend["Budget Month"] == budget_month]
    if month_spend.empty:
        return pd.DataFrame(columns=cols)

    by_cat = month_spend.groupby("AI Category", dropna=False)["Spend Amount"].sum()
    rows: list[dict[str, Any]] = []
    for cat, amt in by_cat.items():
        label = "" if pd.isna(cat) else str(cat)
        avg = float(cat_avg.get(cat, 0.0))
        amount = float(amt)
        vs_avg = amount - avg
        if avg:
            vs_pct = vs_avg / avg * 100
        else:
            vs_pct = 100.0 if amount else 0.0
        pct_month = (amount / month_total * 100) if month_total else 0.0
        rows.append(
            {
                "AI Category": label,
                "Month Spend": round(amount, 2),
                "Avg Monthly Spend": round(avg, 2),
                "vs Avg": round(vs_avg, 2),
                "vs Avg %": round(vs_pct, 1),
                "% of Month": round(pct_month, 1),
            }
        )

    detail = pd.DataFrame(rows)
    if detail.empty:
        return detail
    detail["_dev_sort"] = detail["vs Avg %"].abs()
    detail = detail.sort_values(
        ["_dev_sort", "Month Spend"], ascending=[False, False], kind="stable"
    ).drop(columns=["_dev_sort"])
    return detail.reset_index(drop=True)


def write_spending_insights_sheet(
    writer: pd.ExcelWriter,
    df: pd.DataFrame,
    *,
    ingest_warning_count: int = 0,
) -> dict[str, Any]:
    """
    Summary sheet: monthly expense overview with hyperlinks to per-month category
    breakdown (deviation-first, then spend descending).
    """
    _ = ingest_warning_count  # reserved for future ingest notes on Summary
    spend = _spend_rows(df)
    overview = build_monthly_expense_overview(df, spend)
    sheet = SUMMARY_INSIGHTS_SHEET

    if overview.empty:
        pd.DataFrame(
            [["No included expense rows to build spending insights."]],
            columns=["Note"],
        ).to_excel(writer, sheet_name=sheet, index=False)
        return {"overview_rows": 0, "overview_header_row": 1, "detail_header_rows": [], "month_anchors": {}}

    overview_header_row = 1
    overview.to_excel(writer, sheet_name=sheet, index=False, startrow=0)

    cat_avg = _category_avg_monthly_spend(spend)
    monthly_totals = spend.groupby("Budget Month", dropna=False)["Spend Amount"].sum()
    budget_months = sorted(monthly_totals.index.tolist())

    start_detail = len(overview) + 3
    pd.DataFrame(
        [
            [
                "Category breakdown by month — click View Categories above, or scroll. "
                "Sorted: largest deviation from category average, then spend."
            ]
        ]
    ).to_excel(writer, sheet_name=sheet, index=False, header=False, startrow=start_detail - 1)

    current_row = start_detail + 1
    month_anchors: dict[str, int] = {}
    detail_header_rows: list[int] = []

    for m in budget_months:
        month_key = str(m)
        month_anchors[month_key] = current_row
        pd.DataFrame([[f"── {month_key} ──"]]).to_excel(
            writer,
            sheet_name=sheet,
            index=False,
            header=False,
            startrow=current_row - 1,
        )
        current_row += 1

        detail = build_month_category_breakdown(
            spend,
            month_key,
            cat_avg=cat_avg,
            month_total=float(monthly_totals.get(m, 0.0)),
        )
        detail_header_rows.append(current_row)
        detail.to_excel(writer, sheet_name=sheet, index=False, startrow=current_row - 1)
        current_row += len(detail) + 2

    return {
        "overview_rows": len(overview),
        "overview_header_row": overview_header_row,
        "detail_header_rows": detail_header_rows,
        "month_anchors": month_anchors,
    }


def update_history_workbook(
    df: pd.DataFrame,
    history_path: Path,
    *,
    source_file: str,
    ingest_warning_count: int = 0,
) -> dict[str, int]:
    """
    Merge this run into transaction-history.xlsx (deduped by Transaction ID).
    Returns row counts per sheet written.
    """
    run_export = annotate_history_run(
        prepare_transaction_export_df(df), source_file=source_file
    )
    existing = load_history_workbook(history_path)
    prior_ledger = rebuild_ledger_from_history(existing)

    ledger = merge_dedupe_transactions(prior_ledger, run_export)

    income = ledger[ledger["Flow Type"] == "Income"].copy()
    expenses = sort_expenses_for_export(ledger[ledger["Flow Type"] == "Expense"].copy())
    adjustments = ledger[ledger["Flow Type"] == "Adjustment"].copy()

    ledger_analytics = ledger.copy()
    if "Amount_Numeric" not in ledger_analytics.columns:
        ledger_analytics["Amount_Numeric"] = ledger_analytics["Amount"].apply(parse_amount)

    with open_excel_workbook(history_path) as writer:
        income.to_excel(writer, sheet_name="Income", index=False)
        expenses.to_excel(writer, sheet_name="Expenses", index=False)
        adjustments.to_excel(writer, sheet_name="Adjustments", index=False)
        insights_layout = write_spending_insights_sheet(writer, ledger_analytics)
        baseline_layout = write_baseline_sheet(writer, ledger_analytics)
        for sheet_name in TRANSACTION_SHEETS:
            if sheet_name in writer.sheets:
                format_transaction_worksheet(
                    writer.sheets[sheet_name],
                    hide_descriptions=sheet_name in DETAIL_TRANSACTION_SHEETS,
                    short_date_columns=("Date",)
                    if sheet_name in ("Raw Data", "Income")
                    else (),
                    enable_autofilter=sheet_name == "Expenses",
                )
        if SUMMARY_INSIGHTS_SHEET in writer.sheets:
            format_spending_insights_worksheet(
                writer.sheets[SUMMARY_INSIGHTS_SHEET], **insights_layout
            )
        if "Baseline" in writer.sheets:
            format_baseline_worksheet(writer.sheets["Baseline"], **baseline_layout)

    written = pd.concat([income, expenses, adjustments], ignore_index=True)
    unique_ids = (
        int(written["Transaction ID"].nunique()) if not written.empty else 0
    )
    return {
        "Income": len(income),
        "Expenses": len(expenses),
        "Adjustments": len(adjustments),
        "Summary": insights_layout.get("overview_rows", 0),
        "unique_transactions": unique_ids,
    }


def _norm_lookup_key(category: str, sub_category: str) -> tuple[str, str]:
    return str(category or "").strip().lower(), str(sub_category or "").strip().lower()


def build_merchant_category_map(
    lookups: dict[str, pd.DataFrame],
) -> dict[str, pd.Series]:
    """Merchant Key (lower) -> lookup row. Falls back to legacy Categories sheet."""
    result: dict[str, pd.Series] = {}
    sheet = lookups.get(MERCHANT_CATEGORIES_SHEET)
    if sheet is not None and not sheet.empty and "Merchant Key" in sheet.columns:
        for _, row in sheet.iterrows():
            mk = str(row.get("Merchant Key", "") or "").strip().lower()
            if mk:
                result[mk] = row
        return result

    categories = lookups.get("Categories")
    if categories is not None and not categories.empty:
        for _, row in categories.iterrows():
            mk = str(row.get("AI Sub-Category", "") or "").strip().lower()
            if mk:
                result[mk] = row
    return result


def apply_merchant_category_lookup(
    df: pd.DataFrame,
    lookups: dict[str, pd.DataFrame],
    *,
    spend_mask: pd.Series | None = None,
) -> int:
    """Apply per-merchant categorization from MerchantCategories (or legacy Categories)."""
    by_merchant = build_merchant_category_map(lookups)
    if not by_merchant:
        return 0

    updated = 0
    from_merchant_sheet = (
        lookups.get(MERCHANT_CATEGORIES_SHEET) is not None
        and not lookups.get(MERCHANT_CATEGORIES_SHEET, pd.DataFrame()).empty
    )

    for idx, row in df.iterrows():
        mk = _row_merchant_key(row).lower()
        match = by_merchant.get(mk)
        if match is None:
            continue

        sub = str(match.get("AI Sub-Category", "") or "").strip()
        apply_sub = bool(sub)
        if not from_merchant_sheet and sub.lower() == mk:
            apply_sub = False

        for col, target in (
            ("AI Category", "AI Category"),
            ("Type", "Type"),
            ("Sub-Type", "Sub-Type"),
            ("Flow Type", "Flow Type"),
            ("Classification", "Classification"),
        ):
            val = match.get(col, "")
            if pd.notna(val) and str(val).strip() not in ("", "nan"):
                df.at[idx, target] = str(val).strip()

        if apply_sub and sub.lower() != mk:
            df.at[idx, "AI Sub-Category"] = sub

        if spend_mask is None or bool(spend_mask.loc[idx]):
            tier = match.get("Budget Tier", "")
            if pd.notna(tier) and str(tier).strip() not in ("", "nan", "Review"):
                df.at[idx, "Budget Tier"] = str(tier).strip()
        updated += 1

    return updated


def load_lookup_workbook(lookup_path: Path) -> dict[str, pd.DataFrame]:
    """Load lookup sheets; missing file or sheet returns empty dict entries."""
    sheets: dict[str, pd.DataFrame] = {}
    if not lookup_path.exists():
        return sheets
    try:
        xl = pd.ExcelFile(lookup_path)
    except Exception:
        return sheets
    for name in LOOKUP_SHEETS:
        if name in xl.sheet_names:
            sheets[name] = pd.read_excel(xl, sheet_name=name)
    if CUSTOM_RULES_SHEET in xl.sheet_names:
        sheets[CUSTOM_RULES_SHEET] = pd.read_excel(xl, sheet_name=CUSTOM_RULES_SHEET)
    if EXPENSE_CADENCE_RULES_SHEET in xl.sheet_names:
        sheets[EXPENSE_CADENCE_RULES_SHEET] = pd.read_excel(
            xl, sheet_name=EXPENSE_CADENCE_RULES_SHEET
        )
    return sheets


def apply_lookup_rules(
    df: pd.DataFrame,
    lookups: dict[str, pd.DataFrame],
    *,
    spend_mask: pd.Series | None = None,
) -> int:
    """
    Apply CategoryRules (by source Category) then merchant Categories lookup.
    Budget Tier from lookup is applied only on spend rows (when spend_mask is set).
    Returns approximate count of rows touched by lookup.
    """
    updated = 0
    rules = lookups.get("CategoryRules")
    if rules is not None and not rules.empty and "Source Category" in rules.columns:
        rules = rules.copy()
        rules["Source Category"] = rules["Source Category"].fillna("").astype(str).str.strip()
        rules_by_source = {
            row["Source Category"]: row
            for _, row in rules.iterrows()
            if row["Source Category"]
        }
        for idx, row in df.iterrows():
            src = str(row.get("Category", "") or "").strip()
            rule = rules_by_source.get(src)
            if rule is None:
                continue
            for col, target in (
                ("AI Category", "AI Category"),
                ("Type", "Type"),
                ("Sub-Type", "Sub-Type"),
            ):
                val = rule.get(col, "")
                if pd.notna(val) and str(val).strip() not in ("", "nan"):
                    df.at[idx, target] = str(val).strip()
            if spend_mask is None or bool(spend_mask.loc[idx]):
                tier = rule.get("Budget Tier", "")
                if pd.notna(tier) and str(tier).strip() not in ("", "nan", "Review"):
                    df.at[idx, "Budget Tier"] = str(tier).strip()
            updated += 1

    categories = lookups.get("Categories")
    if categories is None or categories.empty:
        return updated

    categories = categories.copy()
    categories["AI Category"] = categories["AI Category"].fillna("").astype(str)
    categories["AI Sub-Category"] = categories["AI Sub-Category"].fillna("").astype(str)

    lookup_by_pair: dict[tuple[str, str], pd.Series] = {}
    for _, row in categories.iterrows():
        key = _norm_lookup_key(row["AI Category"], row["AI Sub-Category"])
        if not key[1]:
            continue
        lookup_by_pair[key] = row

    for idx, row in df.iterrows():
        if not _is_semantic_sub_category(row):
            continue
        key = _norm_lookup_key(row.get("AI Category", ""), row.get("AI Sub-Category", ""))
        match = lookup_by_pair.get(key)
        if match is None:
            continue
        for col in ("AI Category", "AI Sub-Category", "Type", "Sub-Type"):
            val = match.get(col, "")
            if pd.notna(val) and str(val).strip() not in ("", "nan"):
                df.at[idx, col] = str(val).strip()
        if spend_mask is None or bool(spend_mask.loc[idx]):
            tier = match.get("Budget Tier", "")
            if pd.notna(tier) and str(tier).strip() not in ("", "nan", "Review"):
                df.at[idx, "Budget Tier"] = str(tier).strip()
        updated += 1

    business_rules = lookups.get("BusinessCategoryRules")
    if business_rules is not None and not business_rules.empty:
        if "Generated Description" in business_rules.columns:
            business_rules = business_rules.copy()
            by_desc: dict[str, pd.Series] = {}
            for _, brow in business_rules.iterrows():
                key = str(brow.get("Generated Description", "") or "").strip().lower()
                if key:
                    by_desc[key] = brow
            for idx, row in df.iterrows():
                if not is_business_row(row):
                    continue
                key = str(row.get("Generated Description", "") or "").strip().lower()
                match = by_desc.get(key)
                if match is None:
                    continue
                for col, target in (
                    ("AI Category", "AI Category"),
                    ("AI Sub-Category", "AI Sub-Category"),
                    ("Type", "Type"),
                    ("Sub-Type", "Sub-Type"),
                ):
                    val = match.get(col, "")
                    if pd.notna(val) and str(val).strip() not in ("", "nan"):
                        df.at[idx, target] = str(val).strip()
                if spend_mask is None or bool(spend_mask.loc[idx]):
                    tier = match.get("Budget Tier", "")
                    if pd.notna(tier) and str(tier).strip() not in ("", "nan", "Review"):
                        df.at[idx, "Budget Tier"] = str(tier).strip()
                updated += 1

    return updated


def classify_review_mask(df: pd.DataFrame) -> pd.Series:
    """Spend rows that still need LLM category + semantic sub-category."""
    df = ensure_merchant_key_column(df)
    spend = (
        (df["Flow Type"] == "Expense")
        & (df["Amount_Numeric"] < 0)
        & (df["Include in Spend?"] == "Y")
    )
    sub = df["AI Sub-Category"].fillna("").astype(str).str.strip()
    mk = df["Merchant Key"].fillna("").astype(str).str.strip()
    missing_sub = sub == ""
    legacy_sub = sub.str.lower() == mk.str.lower()
    review_tier = df["Budget Tier"].fillna("").astype(str).str.strip() == "Review"
    return spend & (missing_sub | legacy_sub | review_tier)


def _build_category_rules_from_df(df: pd.DataFrame) -> pd.DataFrame:
    """Default mapping by source Category (one row per source category)."""
    spend = df[(df.get("Include in Spend?", "N") == "Y")].copy()
    if spend.empty or "Category" not in spend.columns:
        return pd.DataFrame(
            columns=["Source Category", "AI Category", "Budget Tier", "Type", "Sub-Type", "Notes"]
        )

    grouped = (
        spend.groupby("Category", dropna=False)
        .agg(
            **{
                "AI Category": ("AI Category", lambda s: s.mode().iat[0] if len(s) else ""),
                "Budget Tier": ("Budget Tier", lambda s: s.mode().iat[0] if len(s) else ""),
                "Type": ("Type", lambda s: s.mode().iat[0] if len(s) else ""),
            }
        )
        .reset_index()
        .rename(columns={"Category": "Source Category"})
    )
    grouped["Sub-Type"] = ""
    grouped["Notes"] = ""
    return grouped.sort_values("Source Category").reset_index(drop=True)


def update_lookup_workbook(
    df: pd.DataFrame,
    lookup_path: Path,
    *,
    client: OpenAI | None = None,
    model: str = "",
    batch_size: int = BATCH_SIZE,
    use_json_mode: bool = False,
    suggested_business: pd.DataFrame | None = None,
    new_description_entries: pd.DataFrame | None = None,
    rebuild_description_lookup: bool = False,
    custom_rules_sheet: pd.DataFrame | None = None,
) -> None:
    """Merge enriched transaction data into the shared lookup workbook."""

    spend_only = df[df.get("Include in Spend?", "N") == "Y"].copy()
    if "Merchant Key" not in spend_only.columns:
        spend_only["Merchant Key"] = spend_only.apply(merchant_key, axis=1)

    semantic_spend = spend_only[
        spend_only.apply(_is_semantic_sub_category, axis=1)
    ].copy()

    merchant_categories = (
        spend_only.groupby("Merchant Key", dropna=False)
        .agg(
            **{
                "AI Category": ("AI Category", lambda s: s.mode().iat[0] if len(s) else ""),
                "AI Sub-Category": (
                    "AI Sub-Category",
                    lambda s: s.mode().iat[0] if len(s) else "",
                ),
                "Budget Tier": ("Budget Tier", lambda s: s.mode().iat[0] if len(s) else ""),
                "Type": ("Type", lambda s: s.mode().iat[0] if len(s) else ""),
                "Transaction Count": ("Merchant Key", "size"),
            }
        )
        .reset_index()
    )
    merchant_categories["Notes"] = ""
    merchant_categories = merchant_categories[
        merchant_categories.apply(_merchant_category_is_semantic, axis=1)
    ].copy()

    categories = (
        semantic_spend.groupby(["AI Category", "AI Sub-Category"], dropna=False)
        .agg(
            **{
                "Transaction Count": ("AI Sub-Category", "size"),
                "Budget Tier": ("Budget Tier", lambda s: s.mode().iat[0] if len(s) else ""),
                "Type": ("Type", lambda s: s.mode().iat[0] if len(s) else ""),
            }
        )
        .reset_index()
    )
    categories["Sub-Type"] = ""
    categories["Notes"] = ""
    categories = categories.sort_values(["AI Category", "AI Sub-Category"]).reset_index(drop=True)

    types = (
        spend_only.groupby(["Type", "Sub-Type"], dropna=False)
        .size()
        .reset_index(name="Transaction Count")
    )
    types["Notes"] = ""

    category_rules = _build_category_rules_from_df(df)

    existing = load_lookup_workbook(lookup_path)
    old_categories = existing.get("Categories")
    old_types = existing.get("Types")
    old_rules = existing.get("CategoryRules")
    old_merchant = existing.get(MERCHANT_CATEGORIES_SHEET)

    if (old_merchant is None or old_merchant.empty) and old_categories is not None:
        legacy_rows: list[dict[str, Any]] = []
        for _, row in old_categories.iterrows():
            mk = str(row.get("AI Sub-Category", "") or "").strip()
            if not mk:
                continue
            legacy_rows.append(
                {
                    "Merchant Key": mk,
                    "AI Category": row.get("AI Category", ""),
                    "AI Sub-Category": "",
                    "Budget Tier": row.get("Budget Tier", ""),
                    "Type": row.get("Type", ""),
                    "Transaction Count": row.get("Transaction Count", ""),
                    "Notes": "migrated from legacy Categories",
                }
            )
        if legacy_rows:
            old_merchant = pd.DataFrame(legacy_rows, columns=list(MERCHANT_CATEGORY_COLUMNS))

    legacy_merchant_keys: set[str] = set()
    if old_categories is not None and not old_categories.empty:
        legacy_merchant_keys = {
            str(v).strip().lower()
            for v in old_categories["AI Sub-Category"].fillna("").astype(str)
            if str(v).strip()
        }

    if old_merchant is not None and not old_merchant.empty:
        mk_key = "Merchant Key"
        old_merchant = old_merchant.copy()
        old_merchant[mk_key] = old_merchant[mk_key].fillna("").astype(str)
        merchant_categories[mk_key] = merchant_categories[mk_key].fillna("").astype(str)
        old_mk = set(old_merchant[mk_key].str.strip().str.lower())
        new_mk_only = merchant_categories[
            ~merchant_categories[mk_key].str.strip().str.lower().isin(old_mk)
        ].copy()
        old_merchant = old_merchant.merge(
            merchant_categories[[mk_key, "AI Sub-Category", "Transaction Count"]],
            on=mk_key,
            how="left",
            suffixes=("", "_new"),
        )
        if "AI Sub-Category_new" in old_merchant.columns:
            has_semantic = old_merchant["AI Sub-Category_new"].fillna("").astype(str).str.strip() != ""
            old_merchant.loc[has_semantic, "AI Sub-Category"] = old_merchant.loc[
                has_semantic, "AI Sub-Category_new"
            ]
            old_merchant.loc[has_semantic, "Transaction Count"] = old_merchant.loc[
                has_semantic, "Transaction Count_new"
            ].fillna(old_merchant.loc[has_semantic, "Transaction Count"])
            old_merchant = old_merchant.drop(
                columns=[c for c in old_merchant.columns if c.endswith("_new")]
            )
        merchant_categories = pd.concat([old_merchant, new_mk_only], ignore_index=True)
        merchant_categories = merchant_categories.sort_values(mk_key).reset_index(drop=True)

    if old_categories is not None and not old_categories.empty:
        old_categories = old_categories.copy()
        old_categories["_sub_lc"] = (
            old_categories["AI Sub-Category"].fillna("").astype(str).str.strip().str.lower()
        )
        old_categories = old_categories[
            ~old_categories["_sub_lc"].isin(legacy_merchant_keys)
        ].drop(columns=["_sub_lc"])

    if old_categories is not None and not old_categories.empty:
        key_cols = ["AI Category", "AI Sub-Category"]
        for col in key_cols:
            old_categories[col] = old_categories[col].fillna("").astype(str)
            categories[col] = categories[col].fillna("").astype(str)
        old_keyset = set(
            tuple(x) for x in old_categories[key_cols].values.tolist()
        )

        new_keyset = set(
            tuple(x) for x in categories[key_cols].values.tolist()
        )

        # Append only truly new pairs; keep existing rows as-is.
        new_only = categories[
            ~categories[key_cols]
            .fillna("")
            .astype(str)
            .apply(lambda r: tuple(r.values.tolist()), axis=1)
            .isin(old_keyset)
        ].copy()

        # Update Transaction Count for existing rows (Budget Tier/Type/Notes preserved).
        old_categories = old_categories.copy()
        computed_counts = categories[key_cols + ["Transaction Count"]].copy()
        old_categories = old_categories.merge(
            computed_counts, on=key_cols, how="left", suffixes=("", "_computed")
        )
        if "Transaction Count_computed" in old_categories.columns:
            old_categories["Transaction Count"] = old_categories[
                "Transaction Count_computed"
            ].fillna(old_categories["Transaction Count"])
            old_categories = old_categories.drop(columns=["Transaction Count_computed"])

        categories = pd.concat([old_categories, new_only], ignore_index=True)
        categories = categories.sort_values(key_cols).reset_index(drop=True)

    if old_types is not None and not old_types.empty:
        type_key_cols = ["Type", "Sub-Type"]
        for col in type_key_cols:
            old_types[col] = old_types[col].fillna("").astype(str)
            types[col] = types[col].fillna("").astype(str)
        old_type_keyset = set(
            tuple(x) for x in old_types[type_key_cols].values.tolist()
        )
        new_only_types = types[
            ~types[type_key_cols]
            .apply(lambda r: tuple(r.values.tolist()), axis=1)
            .isin(old_type_keyset)
        ].copy()

        # Update Transaction Count for existing types only
        old_types = old_types.copy()
        computed_type_counts = types[type_key_cols + ["Transaction Count"]].copy()
        old_types = old_types.merge(
            computed_type_counts, on=type_key_cols, how="left", suffixes=("", "_computed")
        )
        if "Transaction Count_computed" in old_types.columns:
            old_types["Transaction Count"] = old_types[
                "Transaction Count_computed"
            ].fillna(old_types["Transaction Count"])
            old_types = old_types.drop(columns=["Transaction Count_computed"])

        types = pd.concat([old_types, new_only_types], ignore_index=True)
        types = types.sort_values(type_key_cols).reset_index(drop=True)

    if old_rules is not None and not old_rules.empty and "Source Category" in old_rules.columns:
        rule_cols = ["Source Category", "AI Category", "Budget Tier", "Type", "Sub-Type", "Notes"]
        old_rules = old_rules.copy()
        for col in rule_cols:
            if col not in old_rules.columns:
                old_rules[col] = ""
        old_rules["Source Category"] = old_rules["Source Category"].fillna("").astype(str)
        category_rules["Source Category"] = category_rules["Source Category"].fillna("").astype(str)
        old_rule_keys = set(old_rules["Source Category"].tolist())
        new_rules_only = category_rules[
            ~category_rules["Source Category"].isin(old_rule_keys)
        ].copy()
        category_rules = pd.concat([old_rules[rule_cols], new_rules_only], ignore_index=True)
        category_rules = category_rules.sort_values("Source Category").reset_index(drop=True)

    business_rules = existing.get("BusinessCategoryRules")
    if suggested_business is None and client is not None and model:
        suggested_business = enrich_business_lookup_rules(
            df,
            client,
            model,
            existing,
            batch_size=batch_size,
            use_json_mode=use_json_mode,
        )
    business_rules = merge_business_category_rules(
        business_rules, suggested_business if suggested_business is not None else pd.DataFrame()
    )

    description_lookup = merge_description_lookup(
        existing.get("DescriptionLookup"),
        new_description_entries,
        rebuild=rebuild_description_lookup,
    )

    if merchant_categories.empty:
        merchant_categories = pd.DataFrame(columns=list(MERCHANT_CATEGORY_COLUMNS))

    with open_excel_workbook(lookup_path) as writer:
        categories.to_excel(writer, sheet_name="Categories", index=False)
        merchant_categories.to_excel(writer, sheet_name=MERCHANT_CATEGORIES_SHEET, index=False)
        category_rules.to_excel(writer, sheet_name="CategoryRules", index=False)
        types.to_excel(writer, sheet_name="Types", index=False)
        business_rules.to_excel(writer, sheet_name="BusinessCategoryRules", index=False)
        description_lookup.to_excel(writer, sheet_name="DescriptionLookup", index=False)
        custom_out = normalize_custom_rules_sheet(
            custom_rules_sheet if custom_rules_sheet is not None else existing.get(CUSTOM_RULES_SHEET)
        )
        custom_out.to_excel(writer, sheet_name=CUSTOM_RULES_SHEET, index=False)
        cadence_rules = normalize_expense_cadence_rules_sheet(
            existing.get(EXPENSE_CADENCE_RULES_SHEET)
        )
        cadence_rules.to_excel(writer, sheet_name=EXPENSE_CADENCE_RULES_SHEET, index=False)


def load_csv(path: Path) -> pd.DataFrame:
    df = pd.read_csv(
        path,
        dtype=str,
        encoding="utf-8-sig",
        engine="python",
        on_bad_lines="skip",
    )
    df.columns = [c.strip() for c in df.columns]
    if "Amount" in df.columns:
        df["Amount_Numeric"] = df["Amount"].apply(parse_amount)
    return df


def build_transaction_summary(row: pd.Series, index: int) -> dict[str, Any]:
    return {
        "index": index,
        "date": row.get("Date", ""),
        "amount": row.get("Amount", ""),
        "amount_numeric": row.get("Amount_Numeric", parse_amount(row.get("Amount", 0))),
        "original_category": row.get("Category", ""),
        "description": (
            row.get("Generated Description")
            or row.get("Simple Description")
            or row.get("Original Description")
            or ""
        )[:200],
        "account": row.get("Account Name", ""),
    }


def extract_json_payload(text: str) -> dict[str, Any] | list[Any]:
    """Parse JSON from model output, including ```json fenced blocks."""
    text = (text or "").strip()
    if text.startswith("```"):
        lines = text.split("\n")
        if lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        text = "\n".join(lines).strip()
    return json.loads(text)


def _role_model_env(role: str | None) -> str | None:
    """PIPELINE_MODEL or CHAT_MODEL when LM Studio loads multiple models."""
    if not role:
        return None
    key = f"{role.strip().lower()}_model"
    return os.getenv(key.upper(), "").strip() or None


def resolve_provider_config(
    provider_arg: str,
    *,
    base_url_arg: str | None,
    model_arg: str | None,
    role: str | None = None,
) -> tuple[str, str, str]:
    """
    Return (provider, base_url, model) from CLI args and environment.

    Precedence (highest first):
      1. Explicit CLI/API args (base_url_arg, model_arg)
      2. Role env (PIPELINE_MODEL or CHAT_MODEL when role is set)
      3. Generic env overrides (LLM_BASE_URL, LLM_MODEL)
      4. Provider-specific env (LM_STUDIO_*, OLLAMA_*, OPENAI_*)
         chosen by LLM_PROVIDER
    """
    env_provider = os.getenv("LLM_PROVIDER", "").strip().lower()
    provider = provider_arg
    if provider == "auto":
        if env_provider in LOCAL_PROVIDERS | {"openai"}:
            provider = env_provider
        else:
            provider = "lmstudio"

    env_base_url = os.getenv("LLM_BASE_URL", "").strip() or None
    env_model = os.getenv("LLM_MODEL", "").strip() or None
    role_model = _role_model_env(role)
    base_override = base_url_arg or env_base_url
    model_override = model_arg or role_model or env_model

    if provider == "ollama":
        base_url = (base_override or OLLAMA_BASE_URL).rstrip("/")
        model = model_override or OLLAMA_MODEL
    elif provider == "lmstudio":
        base_url = (base_override or LM_STUDIO_BASE_URL).rstrip("/")
        model = model_override or LM_STUDIO_MODEL
    elif provider == "openai":
        base_url = ""
        model = model_override or OPENAI_MODEL
    else:
        raise ValueError(f"Unknown provider: {provider}")

    return provider, base_url, model


def create_client(provider: str, *, base_url: str) -> OpenAI:
    """Return an OpenAI-compatible client for the chosen provider."""
    if provider in LOCAL_PROVIDERS:
        api_key_env = "OLLAMA_API_KEY" if provider == "ollama" else "LM_STUDIO_API_KEY"
        default_key = "ollama" if provider == "ollama" else "lm-studio"
        return OpenAI(
            base_url=base_url.rstrip("/"),
            api_key=os.getenv(api_key_env, default_key),
            timeout=REQUEST_TIMEOUT,
        )
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise ValueError(
            "OPENAI_API_KEY is not set. Use --provider ollama or lmstudio for a local "
            "OpenAI-compatible server, or set OPENAI_API_KEY for OpenAI cloud."
        )
    return OpenAI(api_key=api_key, timeout=REQUEST_TIMEOUT)


def classify_batch(
    client: OpenAI,
    transactions: list[dict[str, Any]],
    model: str,
    *,
    use_json_mode: bool,
) -> list[dict[str, Any]]:
    user_content = json_for_prompt(transactions)
    messages = [
        {"role": "system", "content": CLASSIFICATION_PROMPT},
        {
            "role": "user",
            "content": (
                "Classify these transactions. Respond with a single JSON object "
                '{"results": [ ... ]} where each item has index, section, '
                "category, sub_category, type, sub_type. No markdown, no commentary.\n\n"
                f"Transactions:\n{user_content}"
            ),
        },
    ]
    kwargs: dict[str, Any] = {
        "model": model,
        "temperature": 0.1,
        "messages": messages,
    }
    if use_json_mode:
        kwargs["response_format"] = {"type": "json_object"}

    response = client.chat.completions.create(**kwargs)
    raw = response.choices[0].message.content or "{}"
    try:
        parsed = extract_json_payload(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Could not parse model JSON: {raw[:500]}") from exc

    results = parsed.get("results", parsed if isinstance(parsed, list) else [])
    if not isinstance(results, list):
        raise ValueError(f"Unexpected AI response shape: {raw[:500]}")
    return results


def classify_all(
    df: pd.DataFrame,
    client: OpenAI,
    model: str,
    batch_size: int = BATCH_SIZE,
    *,
    use_json_mode: bool,
) -> pd.DataFrame:
    summaries = [build_transaction_summary(row, i) for i, row in df.iterrows()]
    all_results: dict[int, dict[str, str]] = {}

    for start in range(0, len(summaries), batch_size):
        batch = summaries[start : start + batch_size]
        print(
            f"  Classifying transactions {start + 1}–{start + len(batch)} of {len(summaries)}...",
            flush=True,
        )
        batch_results = classify_batch(client, batch, model, use_json_mode=use_json_mode)
        for item in batch_results:
            idx = int(item.get("index", -1))
            if idx < 0:
                continue
            all_results[idx] = {
                "Section": item.get("section", "Expense"),
                "AI Category": item.get("category", "Other"),
                "AI Sub-Category": item.get("sub_category", ""),
                "Type": item.get("type", "Variable"),
                "Sub-Type": item.get("sub_type", "") or "",
            }

    # Fill any missing rows with sensible defaults from amount sign
    rows = []
    for i in range(len(df)):
        if i in all_results:
            rows.append(all_results[i])
        else:
            amt = summaries[i]["amount_numeric"]
            rows.append(
                {
                    "Section": "Income" if amt > 0 else "Expense",
                    "AI Category": summaries[i]["original_category"] or "Other",
                    "AI Sub-Category": "",
                    "Type": "Variable",
                    "Sub-Type": "",
                }
            )

    enrichment = pd.DataFrame(rows)
    return pd.concat([df.reset_index(drop=True), enrichment], axis=1)


def classify_review_rows(
    df: pd.DataFrame,
    client: OpenAI,
    model: str,
    *,
    batch_size: int,
    use_json_mode: bool,
    on_batch_progress: Callable[[int, int, str], None] | None = None,
) -> pd.DataFrame:
    """Refine AI category and semantic sub-category for spend rows that need it."""
    df = ensure_merchant_key_column(df)
    mask = classify_review_mask(df)
    if not mask.any():
        return df

    review_indices = df[mask].index.tolist()
    print(f"  LLM review rows: {len(review_indices)}", flush=True)

    all_results: dict[int, dict[str, str]] = {}
    summaries = [
        build_transaction_summary(df.loc[i], int(i))  # original index becomes the LLM 'index'
        for i in review_indices
    ]

    total_summaries = len(summaries)
    for start in range(0, total_summaries, batch_size):
        batch = summaries[start : start + batch_size]
        end = start + len(batch)
        batch_msg = f"Classifying review {start + 1}–{end} of {total_summaries}…"
        print(f"  {batch_msg}", flush=True)
        if on_batch_progress:
            on_batch_progress(
                start,
                total_summaries,
                f"{batch_msg} (calling model)",
            )
        try:
            batch_results = classify_batch(
                client, batch, model, use_json_mode=use_json_mode
            )
        except Exception as exc:
            print(f"  LLM failed; using defaults for remaining batches: {exc}", flush=True)
            break
        if on_batch_progress:
            on_batch_progress(end, total_summaries, batch_msg)
        for item in batch_results:
            idx = int(item.get("index", -1))
            if idx < 0:
                continue
            sub = str(item.get("sub_category", "") or "").strip()
            mk = _row_merchant_key(df.loc[idx])
            if sub.lower() == mk.lower():
                sub = ""
            all_results[idx] = {
                "AI Category": item.get("category", df.loc[idx].get("Category", "Other")),
                "AI Sub-Category": sub,
                "Type": item.get("type", df.loc[idx].get("Type", "Variable")),
                "Sub-Type": item.get("sub_type", "") or "",
            }

    # Apply results to df
    df = df.copy()
    for i in review_indices:
        if i in all_results:
            for col, val in all_results[i].items():
                df.at[i, col] = val
        else:
            # Keep defaults
            pass

    return df


def _spend_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Expense rows included in budget spend analysis."""
    mask = (
        (df["Flow Type"] == "Expense")
        & (df.get("Include in Spend?", "N") == "Y")
        & (df["Amount_Numeric"] < 0)
    )
    spend = df.loc[mask].copy()
    if not spend.empty:
        spend["Spend Amount"] = spend["Amount_Numeric"].abs()
    return spend


def build_baseline_sheets(
    df: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """
    Historical spending baseline from included expense rows (by Budget Month).

    Returns overview metrics, per-category baseline, and monthly total spend.
    """
    spend = _spend_rows(df)
    empty_overview = pd.DataFrame(
        columns=["Metric", "Value"],
        data=[["Note", "No included expense rows to build a baseline."]],
    )
    empty_categories = pd.DataFrame(
        columns=[
            "AI Category",
            "Budget Tier",
            "Type",
            "Months With Spend",
            "Irregular Months",
            "Total Spend",
            "Typical Monthly Spend (Core)",
            "Avg Monthly Spend",
            "Min Month",
            "Min Month Spend",
            "Max Month",
            "Max Month Spend",
            "Buffer %",
            "Suggested Monthly Budget",
        ]
    )
    empty_monthly = pd.DataFrame(
        columns=["Budget Month", "Total Spend", "Core Monthly Spend", "vs Monthly Avg", "vs Monthly Avg %"]
    )

    if spend.empty:
        return empty_overview, empty_categories, empty_monthly

    if "In Monthly Run-Rate?" not in spend.columns:
        spend["In Monthly Run-Rate?"] = "Y"
    if "Expense Cadence" not in spend.columns:
        spend["Expense Cadence"] = CADENCE_UNKNOWN

    spend_core = spend[
        spend["In Monthly Run-Rate?"].astype(str).str.strip().str.upper() == "Y"
    ]

    budget_months = sorted(spend["Budget Month"].dropna().unique().tolist())
    n_months = len(budget_months)

    monthly_by_category = (
        spend.groupby(["Budget Month", "AI Category"], dropna=False)["Spend Amount"]
        .sum()
        .reset_index()
    )
    monthly_totals = (
        spend.groupby("Budget Month", dropna=False)["Spend Amount"]
        .sum()
        .sort_index()
    )
    monthly_totals_core = (
        spend_core.groupby("Budget Month", dropna=False)["Spend Amount"]
        .sum()
        .sort_index()
    )
    overall_avg = float(monthly_totals.mean()) if n_months else 0.0
    overall_core_avg = (
        float(monthly_totals_core.mean()) if not monthly_totals_core.empty else 0.0
    )
    overall_min = float(monthly_totals.min()) if n_months else 0.0
    overall_max = float(monthly_totals.max()) if n_months else 0.0
    min_total_month = monthly_totals.idxmin() if n_months else ""
    max_total_month = monthly_totals.idxmax() if n_months else ""

    stability_note = (
        f"Based on {n_months} month(s) in this file."
        if n_months >= BASELINE_RECOMMENDED_MONTHS
        else (
            f"Based on {n_months} month(s); {BASELINE_RECOMMENDED_MONTHS}+ months "
            "recommended for a stable baseline."
        )
    )

    overview = pd.DataFrame(
        [
            {"Metric": "Budget months analyzed", "Value": n_months},
            {"Metric": "Month range", "Value": f"{budget_months[0]} → {budget_months[-1]}"},
            {"Metric": "Overall avg monthly spend", "Value": round(overall_avg, 2)},
            {
                "Metric": "Overall typical core monthly spend",
                "Value": round(overall_core_avg, 2),
            },
            {"Metric": "Lowest month (total spend)", "Value": f"{min_total_month} ({round(overall_min, 2)})"},
            {"Metric": "Highest month (total spend)", "Value": f"{max_total_month} ({round(overall_max, 2)})"},
            {
                "Metric": "Suggested variable buffer",
                "Value": f"{BASELINE_VARIABLE_BUFFER_PCT:g}% on avg (Variable categories)",
            },
            {
                "Metric": "Suggested fixed buffer",
                "Value": f"{BASELINE_FIXED_BUFFER_PCT:g}% on avg (Fixed categories)",
            },
            {"Metric": "Note", "Value": stability_note},
        ]
    )

    category_rows: list[dict[str, Any]] = []
    for category, grp in monthly_by_category.groupby("AI Category", dropna=False):
        by_month = grp.set_index("Budget Month")["Spend Amount"]
        cat_label = "" if pd.isna(category) else str(category)
        cat_spend = spend[spend["AI Category"] == category]
        type_mode = (
            cat_spend["Type"].mode().iat[0]
            if not cat_spend.empty and not cat_spend["Type"].mode().empty
            else "Variable"
        )
        tier_mode = (
            cat_spend["Budget Tier"].mode().iat[0]
            if not cat_spend.empty and not cat_spend["Budget Tier"].mode().empty
            else ""
        )
        buffer_pct = (
            BASELINE_FIXED_BUFFER_PCT
            if str(type_mode).strip().lower() == "fixed"
            else BASELINE_VARIABLE_BUFFER_PCT
        )
        avg_spend = float(by_month.mean())
        cat_core = cat_spend[
            cat_spend["In Monthly Run-Rate?"].astype(str).str.strip().str.upper() == "Y"
        ]
        if not cat_core.empty:
            core_by_month = cat_core.groupby("Budget Month", dropna=False)["Spend Amount"].sum()
            typical_core = float(core_by_month.median())
        else:
            typical_core = 0.0
        irregular_months = 0
        irr = cat_spend[
            cat_spend["Expense Cadence"].isin(
                [CADENCE_YEARLY, CADENCE_ONETIME, CADENCE_UNPLANNED]
            )
        ]
        if not irr.empty and "Budget Month" in irr.columns:
            irregular_months = int(irr["Budget Month"].nunique())
        budget_base = typical_core if typical_core > 0 else avg_spend
        category_rows.append(
            {
                "AI Category": cat_label,
                "Budget Tier": tier_mode,
                "Type": type_mode,
                "Months With Spend": int(by_month.count()),
                "Irregular Months": irregular_months,
                "Total Spend": round(float(by_month.sum()), 2),
                "Typical Monthly Spend (Core)": round(typical_core, 2),
                "Avg Monthly Spend": round(avg_spend, 2),
                "Min Month": by_month.idxmin(),
                "Min Month Spend": round(float(by_month.min()), 2),
                "Max Month": by_month.idxmax(),
                "Max Month Spend": round(float(by_month.max()), 2),
                "Buffer %": buffer_pct,
                "Suggested Monthly Budget": round(budget_base * (1 + buffer_pct / 100), 2),
            }
        )

    categories = pd.DataFrame(category_rows).sort_values(
        "Avg Monthly Spend", ascending=False, na_position="last"
    ).reset_index(drop=True)
    suggested_total = round(float(categories["Suggested Monthly Budget"].sum()), 2)

    monthly_df = monthly_totals.reset_index()
    monthly_df.columns = ["Budget Month", "Total Spend"]
    monthly_df["Total Spend"] = monthly_df["Total Spend"].round(2)
    monthly_df["Core Monthly Spend"] = (
        monthly_df["Budget Month"]
        .map(monthly_totals_core.to_dict())
        .fillna(0)
        .round(2)
    )
    monthly_df["vs Monthly Avg"] = (monthly_df["Total Spend"] - overall_avg).round(2)
    if overall_avg:
        monthly_df["vs Monthly Avg %"] = (
            (monthly_df["Total Spend"] - overall_avg) / overall_avg * 100
        ).round(1)
    else:
        monthly_df["vs Monthly Avg %"] = 0.0

    overview = pd.concat(
        [
            overview,
            pd.DataFrame(
                [{"Metric": "Sum of suggested category budgets", "Value": suggested_total}]
            ),
        ],
        ignore_index=True,
    )

    return overview, categories, monthly_df


def write_baseline_sheet(writer: pd.ExcelWriter, df: pd.DataFrame) -> dict[str, int]:
    """Write Baseline tab: overview + category budget table (monthly drill-down is on Summary)."""
    overview, categories, _monthly = build_baseline_sheets(df)
    sheet = "Baseline"
    overview.to_excel(writer, sheet_name=sheet, index=False, startrow=0)
    start_categories = len(overview) + 2
    categories_header_row = start_categories + 1
    pd.DataFrame(
        [["Category budgets (historical average + buffer — see Summary for month drill-down)"]]
    ).to_excel(
        writer, sheet_name=sheet, index=False, header=False, startrow=start_categories - 1
    )
    categories.to_excel(writer, sheet_name=sheet, index=False, startrow=start_categories)
    return {
        "overview_rows": len(overview),
        "categories_header_row": categories_header_row,
        "monthly_header_row": None,
    }


def write_excel(
    df: pd.DataFrame,
    output_path: Path,
    *,
    ingest_warning_count: int = 0,
    cadence_review: pd.DataFrame | None = None,
) -> None:
    """Write enriched workbook with Raw, Income/Expenses, and Summary tabs."""

    export_df = prepare_transaction_export_df(df)
    raw_sheet = export_df
    income_sheet = export_df[export_df["Flow Type"] == "Income"]
    expense_sheet = sort_expenses_for_export(
        export_df[export_df["Flow Type"] == "Expense"].copy()
    )
    adjustment_sheet = export_df[export_df["Flow Type"] == "Adjustment"]
    with open_excel_workbook(output_path) as writer:
        raw_sheet.to_excel(writer, sheet_name="Raw Data", index=False)
        income_sheet.to_excel(writer, sheet_name="Income", index=False)
        expense_sheet.to_excel(writer, sheet_name="Expenses", index=False)
        adjustment_sheet.to_excel(writer, sheet_name="Adjustments", index=False)
        insights_layout = write_spending_insights_sheet(
            writer, df, ingest_warning_count=ingest_warning_count
        )
        baseline_layout = write_baseline_sheet(writer, df)
        for sheet_name in TRANSACTION_SHEETS:
            if sheet_name in writer.sheets:
                format_transaction_worksheet(
                    writer.sheets[sheet_name],
                    hide_descriptions=sheet_name in DETAIL_TRANSACTION_SHEETS,
                    short_date_columns=("Date",)
                    if sheet_name in ("Raw Data", "Income")
                    else (),
                    enable_autofilter=sheet_name == "Expenses",
                )
        if SUMMARY_INSIGHTS_SHEET in writer.sheets:
            format_spending_insights_worksheet(
                writer.sheets[SUMMARY_INSIGHTS_SHEET], **insights_layout
            )
        if "Baseline" in writer.sheets:
            format_baseline_worksheet(writer.sheets["Baseline"], **baseline_layout)

    print(f"  Raw Data:      {len(raw_sheet)} rows")
    print(f"  Income:        {len(income_sheet)} rows")
    print(f"  Expenses:      {len(expense_sheet)} rows")
    print(f"  Adjustments:   {len(adjustment_sheet)} rows")
    n_months = insights_layout.get("overview_rows", 0)
    if n_months:
        print(f"  Summary:       {n_months} month(s) with category drill-down links")
    _, baseline_categories, _ = build_baseline_sheets(df)
    if not baseline_categories.empty:
        print(f"  Baseline:      {len(baseline_categories)} categories across budget months")


def resolve_paths(input_path: str | None, output_path: str | None) -> tuple[Path, Path]:
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    if input_path:
        inp = Path(input_path)
        if not inp.is_absolute():
            inp = (PROJECT_ROOT / inp).resolve()
    else:
        csv_files = sorted(INPUT_DIR.glob("*.csv"))
        if len(csv_files) == 1:
            inp = csv_files[0]
        elif csv_files:
            inp = csv_files[0]
            print(f"Multiple CSV files in input/; using: {inp.name}")
        else:
            raise FileNotFoundError(
                "No CSV in input/. Pass --input path/to/file.csv"
            )

    if not inp.is_file():
        raise FileNotFoundError(f"Input file not found: {inp}")

    if output_path:
        out = Path(output_path)
        if not out.is_absolute():
            # Bare filename → output/; paths like output/foo.xlsx stay under project root
            if out.parent in (Path("."), Path()):
                out = (OUTPUT_DIR / out.name).resolve()
            else:
                out = (PROJECT_ROOT / out).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
    else:
        out = OUTPUT_DIR / f"{inp.stem}.xlsx"

    return inp, out


def _print_lookup_preflight(lookup_path: Path, *, rebuild_description_lookup: bool) -> None:
    """CLI: summarize lookup workbook contents before the shared pipeline runs."""
    lookups_workbook = load_lookup_workbook(lookup_path)
    if not rebuild_description_lookup:
        description_lookup_map = build_description_lookup_map(lookups_workbook)
        if description_lookup_map:
            print(f"  Loaded {len(description_lookup_map)} description lookup entr(ies)")
    biz_rule_count = len(business_category_rule_keys(lookups_workbook))
    if biz_rule_count:
        print(f"  Loaded {biz_rule_count} business category rule(s)")
    cadence_rules_loaded = normalize_expense_cadence_rules_sheet(
        lookups_workbook.get(EXPENSE_CADENCE_RULES_SHEET) if lookups_workbook else None
    )
    cadence_rule_count = len(
        cadence_rules_loaded[
            cadence_rules_loaded["Generated Description"].fillna("").astype(str).str.strip()
            != ""
        ]
    ) if not cadence_rules_loaded.empty else 0
    if cadence_rule_count:
        print(f"  Loaded {cadence_rule_count} expense cadence rule(s)")


def _print_pipeline_notes(
    result: Any,
    *,
    skip_lookup_update: bool,
    skip_cadence_detection: bool,
    history_path: Path,
) -> None:
    """CLI: post-pipeline messages that used to live inline in main()."""
    from transaction_insight.pipeline import PipelineResult

    if not isinstance(result, PipelineResult):
        return

    stats = result.stats
    new_desc = stats.get("new_description_entries", 0)
    if new_desc:
        print(f"  {new_desc} new description lookup entr(ies) saved")

    if stats.get("business_marked"):
        print(
            f"  Marked {stats['business_marked']} row(s) as Business "
            f"(Generated Description in BusinessCategoryRules)",
            flush=True,
        )
    if stats.get("lookup_rows_touched"):
        print(f"  Applied shared lookup to {stats['lookup_rows_touched']} row(s)")
    if stats.get("business_rules_applied"):
        print(f"  Applied business lookup rules to {stats['business_rules_applied']} row(s)")

    if not skip_lookup_update and not result.custom_rules_sheet.empty:
        active_custom_rules = load_active_custom_rules(result.custom_rules_sheet)
        if not active_custom_rules:
            has_rules = result.custom_rules_sheet["Rule"].fillna("").astype(str).str.strip() != ""
            if has_rules.any():
                print(
                    "  Custom rules present but none Active with valid Compiled Rule; "
                    "set Status=Pending to compile or fix Last Error",
                    flush=True,
                )

    if stats.get("cadence_lookup"):
        print(f"  Applied {stats['cadence_lookup']} expense cadence lookup rule(s)", flush=True)
    if stats.get("cadence_detected"):
        print(
            f"  Cadence detection: {stats['cadence_detected']} row(s) tagged from history patterns",
            flush=True,
        )
    elif (
        not skip_cadence_detection
        and not stats.get("cadence_detected")
        and history_path.exists()
    ):
        print(
            f"  Cadence detection: need {CADENCE_MIN_HISTORY_MONTHS}+ budget months "
            f"in history (set rules in {EXPENSE_CADENCE_RULES_SHEET} meanwhile)",
            flush=True,
        )

    if not result.cadence_review.empty:
        print(
            f"  {len(result.cadence_review)} auto-tagged irregular expense(s) "
            "(see Summary category breakdown vs Avg)",
            flush=True,
        )
    if stats.get("custom_rules_applied"):
        active_count = stats.get("active_custom_rules", stats["custom_rules_applied"])
        print(
            f"  Applied {stats['custom_rules_applied']} custom rule assignment(s) "
            f"from {active_count} active rule(s) (final pass)",
            flush=True,
        )

    hist = stats.get("history")
    if hist:
        print(
            f"  History: {hist['unique_transactions']} unique transaction(s) "
            f"(Income {hist['Income']}, Expenses {hist['Expenses']}, "
            f"Adjustments {hist['Adjustments']})"
        )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert transaction CSV to categorized Excel workbook using AI."
    )
    parser.add_argument(
        "--input", "-i",
        help="Path to input CSV (default: single .csv in input/)",
    )
    parser.add_argument(
        "--output", "-o",
        help="Path to output Excel file (default: output/<csv-stem>.xlsx)",
    )
    parser.add_argument(
        "--provider",
        choices=["auto", "lmstudio", "ollama", "openai"],
        default="auto",
        help="LLM backend: lmstudio, ollama (local), openai (cloud), or auto (uses LLM_PROVIDER from .env)",
    )
    parser.add_argument(
        "--base-url",
        default=None,
        help="OpenAI-compatible base URL for lmstudio/ollama (provider-specific default if omitted)",
    )
    parser.add_argument(
        "--model",
        default=None,
        help="Model id (provider-specific default from .env if omitted)",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=None,
        help=(
            f"Transactions per API call (default: {LOCAL_BATCH_SIZE} for local LLM, "
            f"{BATCH_SIZE} for OpenAI; override with DESCRIPTION_BATCH_SIZE / "
            "CLASSIFICATION_BATCH_SIZE in .env)"
        ),
    )
    parser.add_argument(
        "--skip-lookup-update",
        action="store_true",
        help="Do not read or write the shared lookup file (transaction-lookups.xlsx)",
    )
    parser.add_argument(
        "--rebuild-description-lookup",
        action="store_true",
        help=(
            "Ignore cached descriptions and re-run the description LLM for this file; "
            "matching Source Keys in DescriptionLookup are overwritten on save"
        ),
    )
    parser.add_argument(
        "--lookup-file",
        default=None,
        help=f"Path to shared lookup workbook (default: scripts/{LOOKUP_FILENAME})",
    )
    parser.add_argument(
        "--update-history",
        action="store_true",
        help=(
            f"Merge this run into cumulative history workbook ({HISTORY_FILENAME}). "
            "Off by default so accidental runs do not change history."
        ),
    )
    parser.add_argument(
        "--history-file",
        default=None,
        help=f"Path to cumulative history workbook (default: scripts/{HISTORY_FILENAME}; only with --update-history)",
    )
    parser.add_argument(
        "--skip-cadence-detection",
        action="store_true",
        help=(
            "Do not auto-detect yearly/one-time/unplanned cadence from transaction-history "
            "(lookup rules still apply)"
        ),
    )
    args = parser.parse_args()

    if args.rebuild_description_lookup and args.skip_lookup_update:
        print(
            "Error: --rebuild-description-lookup requires lookup read/write; "
            "do not combine with --skip-lookup-update",
            file=sys.stderr,
        )
        return 1

    try:
        provider, base_url, model = resolve_provider_config(
            args.provider,
            base_url_arg=args.base_url,
            model_arg=args.model,
            role="pipeline",
        )
    except ValueError as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1

    _, description_batch_size, classification_batch_size = resolve_batch_sizes(
        provider, batch_size_arg=args.batch_size
    )

    try:
        input_path, output_path = resolve_paths(args.input, args.output)
    except FileNotFoundError as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1

    lookup_path = Path(args.lookup_file) if args.lookup_file else default_lookup_path()
    history_path = Path(args.history_file) if args.history_file else default_history_path()

    print(f"Input:    {input_path}")
    print(f"Output:   {output_path}")
    print(f"Provider: {provider}")
    print(f"Model:    {model}")
    if provider in LOCAL_PROVIDERS:
        print(f"Base URL: {base_url}")
        print(
            f"Batches:  descriptions={description_batch_size}, "
            f"classification={classification_batch_size}"
        )
    if not args.skip_lookup_update:
        print(f"Lookup:   {lookup_path}" + (" (will create)" if not lookup_path.exists() else ""))
        if args.rebuild_description_lookup:
            print("  Description lookup rebuild: cache ignored, LLM refresh enabled")
    if args.update_history:
        print(f"History:  {history_path}" + (" (will create)" if not history_path.exists() else ""))

    from transaction_insight.config import PipelineConfig
    from transaction_insight.pipeline import run_pipeline

    timer = PhaseTimer()

    with timer.phase("Load CSV"):
        df = load_csv(input_path)
    print(f"Loaded {len(df)} transactions")

    if not args.skip_lookup_update:
        _print_lookup_preflight(lookup_path, rebuild_description_lookup=args.rebuild_description_lookup)

    def on_progress(ev: dict[str, Any]) -> None:
        if ev.get("type") == "phase":
            msg = ev.get("message", "")
            if msg:
                print(msg, flush=True)

    config = PipelineConfig(
        lookup_path=lookup_path,
        history_path=history_path,
        skip_lookup_update=args.skip_lookup_update,
        rebuild_description_lookup=args.rebuild_description_lookup,
        skip_cadence_detection=args.skip_cadence_detection,
        provider=args.provider,
        base_url=args.base_url,
        model=args.model,
        batch_size=args.batch_size,
        source_file=input_path.name,
        update_lookup_workbook=not args.skip_lookup_update,
        update_history=args.update_history,
    )

    result = run_pipeline(
        df, config, on_progress=on_progress, input_path=input_path, timer=timer
    )

    _print_pipeline_notes(
        result,
        skip_lookup_update=args.skip_lookup_update,
        skip_cadence_detection=args.skip_cadence_detection,
        history_path=history_path,
    )

    with timer.phase("Write Excel workbook"):
        print("Writing Excel workbook...")
        write_excel(
            result.dataframe,
            output_path,
            ingest_warning_count=result.ingest_warning_count,
            cadence_review=result.cadence_review,
        )
    print_business_tagging_summary(result.dataframe)
    print(f"Done: {output_path}")

    from transaction_insight.inbox_archive import move_csv_to_processed

    inbox_dir = Path(os.getenv("FINANCE_INBOX_DIR", str(INPUT_DIR)))
    archived_to = move_csv_to_processed(input_path, inbox_dir=inbox_dir)
    if archived_to:
        print(f"Archived CSV: {archived_to}")

    timer.print_summary()
    return 0
